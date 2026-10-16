"""Tests for Issue #35: Multi-Drawing Scope, Revision Management, and Project-Level Takeoff."""

from __future__ import annotations

import io

import openpyxl
from eai_api.app import app
from eai_api.models import SheetRole
from fastapi.testclient import TestClient

client = TestClient(app)


def test_multi_drawing_seed_creates_4_drawings_with_correct_scopes() -> None:
    res = client.post("/api/demo/seed-multi-drawing?reset=true")
    assert res.status_code == 200, res.text
    data = res.json()
    assert data["status"] == "seeded"
    proj_id = data["project_id"]
    drawings = data["drawings"]
    assert len(drawings) == 4

    # Check scopes
    scopes_res = client.get(f"/api/projects/{proj_id}/scopes")
    assert scopes_res.status_code == 200
    scopes = scopes_res.json()
    assert len(scopes) == 4

    f1_r0 = next((s for s in scopes if s["revision"] == "REV0" and s["sheet_key"] == "1F_LIGHT"), None)
    f2_r0 = next((s for s in scopes if s["revision"] == "REV0" and s["sheet_key"] == "2F_LIGHT"), None)
    f1_r1 = next((s for s in scopes if s["revision"] == "REV1" and s["sheet_key"] == "1F_LIGHT"), None)
    f1_ref = next((s for s in scopes if s["sheet_role"] == SheetRole.REFERENCE), None)

    assert f1_r0 is not None
    assert f1_r0["active_for_takeoff"] is True
    assert f1_r0["sheet_role"] == SheetRole.INSTALLATION

    assert f2_r0 is not None
    assert f2_r0["active_for_takeoff"] is True

    assert f1_r1 is not None
    assert f1_r1["active_for_takeoff"] is False  # Inactive initially

    assert f1_ref is not None
    assert f1_ref["active_for_takeoff"] is False
    assert f1_ref["sheet_role"] == SheetRole.REFERENCE


def test_initial_project_takeoff_yields_exact_14_lights() -> None:
    seed_res = client.post("/api/demo/seed-multi-drawing?reset=true")
    assert seed_res.status_code == 200
    proj_id = seed_res.json()["project_id"]

    # Calculate project-wide consolidated takeoff
    takeoff_res = client.get(f"/api/projects/{proj_id}/takeoff")
    assert takeoff_res.status_code == 200
    data = takeoff_res.json()

    assert data["total_active_drawings"] == 2
    assert data["total_inactive_drawings"] == 2

    # Check consolidated summary
    summary = {row["item_code"]: row for row in data["summary"]}
    assert "LIGHT" in summary

    # Ground Truth: 1F REV0 (6) + 2F REV0 (8) = 14 LIGHT EA
    assert summary["LIGHT"]["total_quantity"] == 14
    assert summary["LIGHT"]["unit"] == "EA"
    assert summary["LIGHT"]["active_drawings_count"] == 2

    # OUTLET: 4 + 4 = 8 EA
    assert summary["OUTLET"]["total_quantity"] == 8

    # SWITCH: 3 + 3 = 6 EA
    assert summary["SWITCH"]["total_quantity"] == 6

    # Verify breakdown contains the 2 active drawings
    breakdown = data["drawings_breakdown"]
    assert len(breakdown) == 2
    f1_item = next((b for b in breakdown if b["sheet_key"] == "1F_LIGHT"), None)
    f2_item = next((b for b in breakdown if b["sheet_key"] == "2F_LIGHT"), None)
    assert f1_item is not None
    assert f2_item is not None

    f1_light = next((i for i in f1_item["items"] if i["item_code"] == "LIGHT"), None)
    f2_light = next((i for i in f2_item["items"] if i["item_code"] == "LIGHT"), None)
    assert f1_light is not None
    assert f2_light is not None
    assert f1_light["quantity"] == 6
    assert f2_light["quantity"] == 8


def test_activate_revision_switches_1f_to_rev1_yielding_18_lights() -> None:
    seed_res = client.post("/api/demo/seed-multi-drawing?reset=true")
    proj_id = seed_res.json()["project_id"]

    scopes = client.get(f"/api/projects/{proj_id}/scopes").json()
    f1_r0 = next(s for s in scopes if s["revision"] == "REV0" and s["sheet_key"] == "1F_LIGHT")
    f1_r1 = next(s for s in scopes if s["revision"] == "REV1" and s["sheet_key"] == "1F_LIGHT")

    # Switch active revision from REV0 to REV1
    act_res = client.post(
        f"/api/projects/{proj_id}/drawings/{f1_r1['drawing_id']}/activate-revision",
        json={"active": True},
    )
    assert act_res.status_code == 200
    assert act_res.json()["active_for_takeoff"] is True

    # Re-check scopes: REV0 must be deactivated
    updated_scopes = client.get(f"/api/projects/{proj_id}/scopes").json()
    f1_r0_up = next(s for s in updated_scopes if s["drawing_id"] == f1_r0["drawing_id"])
    f1_r1_up = next(s for s in updated_scopes if s["drawing_id"] == f1_r1["drawing_id"])
    assert f1_r0_up["active_for_takeoff"] is False
    assert f1_r1_up["active_for_takeoff"] is True

    # Project takeoff: 1F REV1 (10) + 2F REV0 (8) = 18 LIGHT EA (+4 delta)
    takeoff_res = client.get(f"/api/projects/{proj_id}/takeoff")
    assert takeoff_res.status_code == 200
    data = takeoff_res.json()

    summary = {row["item_code"]: row for row in data["summary"]}
    assert summary["LIGHT"]["total_quantity"] == 18
    assert summary["LIGHT"]["active_drawings_count"] == 2


def test_revision_comparison_diff_endpoint() -> None:
    seed_res = client.post("/api/demo/seed-multi-drawing?reset=true")
    proj_id = seed_res.json()["project_id"]

    scopes = client.get(f"/api/projects/{proj_id}/scopes").json()
    f1_r0 = next(s for s in scopes if s["revision"] == "REV0" and s["sheet_key"] == "1F_LIGHT")
    f1_r1 = next(s for s in scopes if s["revision"] == "REV1" and s["sheet_key"] == "1F_LIGHT")

    # Compare REV0 vs REV1
    cmp_res = client.get(
        f"/api/projects/{proj_id}/revisions/compare?base_drawing_id={f1_r0['drawing_id']}&compare_drawing_id={f1_r1['drawing_id']}"
    )
    assert cmp_res.status_code == 200
    data = cmp_res.json()

    # Total base = 13 (6 LIGHT + 4 OUTLET + 3 SWITCH)
    # Total compare = 17 (10 LIGHT + 4 OUTLET + 3 SWITCH)
    # Total delta = +4
    assert data["total_base_qty"] == 13
    assert data["total_compare_qty"] == 17
    assert data["total_delta"] == 4

    items_delta = {row["item_code"]: row for row in data["items_delta"]}
    assert "LIGHT" in items_delta
    assert items_delta["LIGHT"]["base_quantity"] == 6
    assert items_delta["LIGHT"]["compare_quantity"] == 10
    assert items_delta["LIGHT"]["delta"] == 4
    assert items_delta["LIGHT"]["change_type"] == "INCREASED"

    # OUTLET and SWITCH remain unchanged
    assert items_delta["OUTLET"]["delta"] == 0
    assert items_delta["OUTLET"]["change_type"] == "UNCHANGED"
    assert items_delta["SWITCH"]["delta"] == 0
    assert items_delta["SWITCH"]["change_type"] == "UNCHANGED"

    # Object deltas: 6 matched, 4 added
    obj_deltas = data["object_deltas"]
    matched = [o for o in obj_deltas if o["match_status"] == "MATCHED"]
    added = [o for o in obj_deltas if o["match_status"] == "ADDED"]
    assert len(matched) >= 6
    assert len(added) == 4


def test_scope_conflict_warning_when_multiple_active() -> None:
    seed_res = client.post("/api/demo/seed-multi-drawing?reset=true")
    proj_id = seed_res.json()["project_id"]

    scopes = client.get(f"/api/projects/{proj_id}/scopes").json()
    f1_r1 = next(s for s in scopes if s["revision"] == "REV1" and s["sheet_key"] == "1F_LIGHT")

    # Forcefully update REV1 to active without deactivating REV0 (manual update)
    client.put(
        f"/api/projects/{proj_id}/drawings/{f1_r1['drawing_id']}/scope",
        json={"active_for_takeoff": True},
    )

    takeoff_res = client.get(f"/api/projects/{proj_id}/takeoff")
    assert takeoff_res.status_code == 200
    data = takeoff_res.json()
    assert len(data["scope_warnings"]) >= 1
    assert "1F_LIGHT" in data["scope_warnings"][0]


def test_comprehensive_4_sheet_project_xlsx_export() -> None:
    seed_res = client.post("/api/demo/seed-multi-drawing?reset=true")
    proj_id = seed_res.json()["project_id"]

    # Export project XLSX
    xlsx_res = client.get(f"/api/projects/{proj_id}/export/xlsx")
    assert xlsx_res.status_code == 200
    assert (
        xlsx_res.headers["Content-Type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_res.content))
    sheet_names = wb.sheetnames
    assert "Project Summary" in sheet_names
    assert "Drawing Breakdown" in sheet_names
    assert "Revision Delta" in sheet_names
    assert "Evidence" in sheet_names

    # Check Sheet 1: Project Summary
    ws_sum = wb["Project Summary"]
    headers = [cell.value for cell in ws_sum[1]]
    assert "총 산출 수량 (total_qty)" in headers
    # Verify LIGHT row has 14
    rows = list(ws_sum.iter_rows(values_only=True))[1:]
    light_row = next((r for r in rows if r[1] == "LIGHT"), None)
    assert light_row is not None
    assert light_row[5] == 14

    # Check Sheet 3: Revision Delta
    ws_delta = wb["Revision Delta"]
    delta_rows = list(ws_delta.iter_rows(values_only=True))[1:]
    delta_light = next((r for r in delta_rows if r[4] == "LIGHT"), None)
    assert delta_light is not None
    assert delta_light[8] == 6   # base_qty
    assert delta_light[9] == 10  # compare_qty
    assert delta_light[10] == 4  # delta

    # Formula injection verification: zero formula cells
    for s_name in sheet_names:
        ws = wb[s_name]
        for row in ws.iter_rows():
            for c in row:
                assert c.data_type != "f", f"Formula vulnerability in {s_name} {c.coordinate}"


def test_project_csv_exports() -> None:
    seed_res = client.post("/api/demo/seed-multi-drawing?reset=true")
    proj_id = seed_res.json()["project_id"]

    # 1. Summary CSV
    csv_sum = client.get(f"/api/projects/{proj_id}/export/csv?type=summary")
    assert csv_sum.status_code == 200
    text = csv_sum.text
    assert "LIGHT" in text
    assert "14" in text

    # 2. Breakdown CSV
    csv_bd = client.get(f"/api/projects/{proj_id}/export/csv?type=breakdown")
    assert csv_bd.status_code == 200
    assert "1F_LIGHT" in csv_bd.text
    assert "2F_LIGHT" in csv_bd.text

    # 3. Delta CSV
    csv_delta = client.get(f"/api/projects/{proj_id}/export/csv?type=delta")
    assert csv_delta.status_code == 200
    assert "REV0 -> REV1" in csv_delta.text or "REV0" in csv_delta.text
    assert "LIGHT" in csv_delta.text

