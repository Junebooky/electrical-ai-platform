"""Comprehensive tests for Mapping Templates, openpyxl cell-level Excel verification, and Demo Seed (Issue #23).

Verifies:
1. Built-in and custom Mapping Template CRUD (list, create, get, delete).
2. Merging vs Replacing mapping templates on drawings.
3. openpyxl cell-level Excel export verification (sheet names, column headers, exact quantity cells).
4. Full live demo project seeding (DXF, PDF, Image 3-format end-to-end integration).
"""

from __future__ import annotations

import io
from pathlib import Path

import openpyxl
import pytest
from eai_api.app import app
from eai_api.storage import Storage
from fastapi.testclient import TestClient

DXF_FIXTURE_PATH = Path("fixtures/dxf/basic_devices.dxf")


@pytest.fixture
def client(tmp_path: Path) -> TestClient:
    from eai_api import app as app_module

    app_module.storage = Storage(root=tmp_path)
    return TestClient(app)


def test_list_builtin_templates(client: TestClient) -> None:
    """Verify built-in mapping templates are available out-of-the-box."""
    resp = client.get("/api/mapping-templates")
    assert resp.status_code == 200
    templates = resp.json()["templates"]

    tpl_ids = {t["id"] for t in templates}
    assert "tpl_standard_elec" in tpl_ids
    assert "tpl_kec_standard" in tpl_ids

    std_tpl = next(t for t in templates if t["id"] == "tpl_standard_elec")
    assert std_tpl["is_builtin"] is True
    assert len(std_tpl["rules"]) >= 8

    # Verify standard mappings: LT-DL -> LIGHT, OT-2P -> OUTLET, SW-1 -> SWITCH
    rule_map = {r["block_name"]: r["item_code"] for r in std_tpl["rules"]}
    assert rule_map.get("LT-DL") == "LIGHT"
    assert rule_map.get("OT-2P") == "OUTLET"
    assert rule_map.get("SW-1") == "SWITCH"


def test_create_and_delete_custom_template(client: TestClient) -> None:
    """Verify creating, retrieving, and deleting a custom mapping template."""
    # 1. Create custom template
    payload = {
        "name": "오피스 맞춤 전기 템플릿",
        "description": "사무실 전용 조명 및 바닥 콘센트 매핑 규칙",
        "rules": [
            {"block_name": "OFFICE-LT", "item_code": "LIGHT", "item_name": "사무실 조명", "unit": "EA"},
            {"block_name": "FL-CON", "item_code": "OUTLET", "item_name": "플로어 콘센트", "unit": "EA"},
        ],
    }
    c_resp = client.post("/api/mapping-templates", json=payload)
    assert c_resp.status_code == 200
    created = c_resp.json()
    tpl_id = created["id"]
    assert created["name"] == "오피스 맞춤 전기 템플릿"
    assert created["is_builtin"] is False
    assert len(created["rules"]) == 2

    # 2. Retrieve template by ID
    g_resp = client.get(f"/api/mapping-templates/{tpl_id}")
    assert g_resp.status_code == 200
    assert g_resp.json()["id"] == tpl_id

    # 3. Built-in templates cannot be deleted
    del_builtin = client.delete("/api/mapping-templates/tpl_standard_elec")
    assert del_builtin.status_code == 400

    # 4. Custom template can be deleted
    del_resp = client.delete(f"/api/mapping-templates/{tpl_id}")
    assert del_resp.status_code == 200
    assert del_resp.json()["status"] == "deleted"

    # Confirm it is gone
    assert client.get(f"/api/mapping-templates/{tpl_id}").status_code == 404


def test_apply_template_merge_and_replace(client: TestClient) -> None:
    """Verify applying templates to drawing with MERGE vs REPLACE strategies.

    Specifically verifies P0-QTY-02: MERGE must preserve existing user-defined rules
    (e.g., LT-DL -> CUSTOM_EM_LIGHT) and only append missing template rules.
    """
    p_resp = client.post("/api/projects", json={"name": "Template Apply Test"})
    proj_id = p_resp.json()["id"]

    dxf_bytes = DXF_FIXTURE_PATH.read_bytes()
    up_resp = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("basic_devices.dxf", dxf_bytes, "application/dxf")},
    )
    dwg_id = up_resp.json()["drawing_id"]

    # 1. Initial manual rules: one non-overlapping block, and one customized block (LT-DL)
    client.put(
        f"/api/projects/{proj_id}/drawings/{dwg_id}/mappings",
        json=[
            {"block_name": "MY_CUSTOM_BLOCK", "item_code": "PANEL", "item_name": "특수 분전반", "unit": "EA"},
            {"block_name": "LT-DL", "item_code": "CUSTOM_EM_LIGHT", "item_name": "사용자 커스텀 비상등", "unit": "EA"},
        ],
    )

    # 2. Apply template with MERGE: must NOT overwrite user's LT-DL mapping!
    merge_resp = client.post(
        f"/api/projects/{proj_id}/drawings/{dwg_id}/apply-template",
        json={"template_id": "tpl_standard_elec", "merge_strategy": "MERGE"},
    )
    assert merge_resp.status_code == 200
    merged_rules = merge_resp.json()
    rule_map = {r["block_name"]: r for r in merged_rules}

    # Preserved user mappings
    assert "MY_CUSTOM_BLOCK" in rule_map
    assert rule_map["MY_CUSTOM_BLOCK"]["item_code"] == "PANEL"
    assert rule_map["LT-DL"]["item_code"] == "CUSTOM_EM_LIGHT"  # P0-QTY-02: NOT overwritten to "LIGHT"
    assert rule_map["LT-DL"]["item_name"] == "사용자 커스텀 비상등"

    # Newly appended template rules
    assert "OT-2P" in rule_map
    assert rule_map["OT-2P"]["item_code"] == "OUTLET"
    assert "SW-1" in rule_map
    assert rule_map["SW-1"]["item_code"] == "SWITCH"

    # 3. Apply template with REPLACE: completely overwrites all rules
    replace_resp = client.post(
        f"/api/projects/{proj_id}/drawings/{dwg_id}/apply-template",
        json={"template_id": "tpl_standard_elec", "merge_strategy": "REPLACE"},
    )
    assert replace_resp.status_code == 200
    replaced_rules = replace_resp.json()
    rep_map = {r["block_name"]: r for r in replaced_rules}
    assert "MY_CUSTOM_BLOCK" not in rep_map
    assert rep_map["LT-DL"]["item_code"] == "LIGHT"  # Now reset to standard template rule


def test_openpyxl_exact_cell_value_verification(client: TestClient) -> None:
    """Verify actual cell values in exported XLSX workbook using openpyxl.

    Directly verifies A01's QA requirement: exact quantity cell values (e.g. 6 EA) and sheet integrity.
    """
    p_resp = client.post("/api/projects", json={"name": "Openpyxl XLSX Test"})
    proj_id = p_resp.json()["id"]

    dxf_bytes = DXF_FIXTURE_PATH.read_bytes()
    up_resp = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("basic_devices.dxf", dxf_bytes, "application/dxf")},
    )
    dwg_id = up_resp.json()["drawing_id"]

    # 1. Apply standard template (maps LT-DL to LIGHT, OT-2P to OUTLET, SW-1 to SWITCH)
    client.post(
        f"/api/projects/{proj_id}/drawings/{dwg_id}/apply-template",
        json={"template_id": "tpl_standard_elec", "merge_strategy": "REPLACE"},
    )

    # 2. Run detection and accept all candidates
    client.post(f"/api/projects/{proj_id}/drawings/{dwg_id}/candidates/detect")
    client.put(
        f"/api/projects/{proj_id}/drawings/{dwg_id}/candidates/review",
        json={"action": "ACCEPT_ALL"},
    )

    # 3. Download XLSX binary
    xlsx_resp = client.get(f"/api/projects/{proj_id}/drawings/{dwg_id}/export/xlsx")
    assert xlsx_resp.status_code == 200
    xlsx_bytes = xlsx_resp.content

    # 4. Parse workbook with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert "Takeoff Summary" in wb.sheetnames
    assert "Takeoff Evidence" in wb.sheetnames

    # 5. Check 'Takeoff Summary' sheet cell values
    ws_summary = wb["Takeoff Summary"]
    headers = [str(cell.value) for cell in ws_summary[1]]
    assert any("item_code" in h for h in headers)
    assert any("quantity" in h for h in headers)
    assert any("unit" in h for h in headers)

    code_col = next(i + 1 for i, h in enumerate(headers) if "item_code" in h)
    qty_col = next(i + 1 for i, h in enumerate(headers) if "quantity" in h)
    unit_col = next(i + 1 for i, h in enumerate(headers) if "unit" in h)

    summary_data: dict[str, int] = {}
    for row in range(2, ws_summary.max_row + 1):
        code_val = ws_summary.cell(row=row, column=code_col).value
        qty_val = ws_summary.cell(row=row, column=qty_col).value
        unit_val = ws_summary.cell(row=row, column=unit_col).value
        if code_val:
            summary_data[str(code_val)] = int(qty_val)
            assert unit_val == "EA"

    # In basic_devices.dxf: 6 LT-DL -> LIGHT: 6 EA, 4 OT-2P -> OUTLET: 4 EA, 3 SW-1 -> SWITCH: 3 EA
    assert summary_data.get("LIGHT") == 6, f"Expected LIGHT to be 6 in XLSX cell, got {summary_data.get('LIGHT')}"
    assert summary_data.get("OUTLET") == 4, f"Expected OUTLET to be 4 in XLSX cell, got {summary_data.get('OUTLET')}"
    assert summary_data.get("SWITCH") == 3, f"Expected SWITCH to be 3 in XLSX cell, got {summary_data.get('SWITCH')}"

    # 6. Check 'Takeoff Evidence' sheet cell values
    ws_evidence = wb["Takeoff Evidence"]
    ev_headers = [str(cell.value) for cell in ws_evidence[1]]
    assert any("source_object_id" in h for h in ev_headers)
    assert any("inclusion_status" in h for h in ev_headers)

    inc_col = next(i + 1 for i, h in enumerate(ev_headers) if "inclusion_status" in h)
    included_count = 0
    for row in range(2, ws_evidence.max_row + 1):
        inc_val = ws_evidence.cell(row=row, column=inc_col).value
        if inc_val == "INCLUDED":
            included_count += 1

    # Total included rows must match total EA count: 6 + 4 + 3 = 13!
    assert included_count == 13, f"Expected 13 included evidence rows in XLSX, got {included_count}"


def test_seed_demo_project_workflow(client: TestClient) -> None:
    """Verify live demo project seeding with DXF, PDF, and Image drawings."""
    resp = client.post("/api/demo/seed")
    assert resp.status_code == 200
    data = resp.json()

    assert data["status"] == "seeded"
    proj_id = data["project_id"]
    drawings = data["drawings"]
    assert len(drawings) == 3

    formats = {d["format"] for d in drawings}
    assert formats == {"DXF", "PDF", "IMAGE"}

    # Verify each seeded drawing has candidates and takeoff
    for d in drawings:
        dwg_id = d["id"]
        cands_resp = client.get(f"/api/projects/{proj_id}/drawings/{dwg_id}/candidates")
        assert cands_resp.status_code == 200
        assert cands_resp.json()["summary"]["total"] > 0

        t_resp = client.get(f"/api/projects/{proj_id}/drawings/{dwg_id}/takeoff")
        assert t_resp.status_code == 200
        takeoff = t_resp.json()
        assert "summary" in takeoff
        assert "evidence" in takeoff


def test_apply_template_unit_validation_on_pdf_and_image(client: TestClient) -> None:
    """Verify P0-MUNIT-03: On PDF and Image drawings, non-EA units are rejected both in save_mappings and apply_template.

    Fail-closed: if rejected with 400, existing mappings must remain completely intact.
    """
    p_resp = client.post("/api/projects", json={"name": "Unit Validation Project"})
    proj_id = p_resp.json()["id"]

    # 1. Create a template with non-EA unit (e.g. M)
    c_tpl = client.post(
        "/api/mapping-templates",
        json={
            "name": "비EA 단위 템플릿",
            "description": "케이블 길이 산출용 미터 단위 템플릿",
            "rules": [{"block_name": "CABLE-M", "item_code": "CABLE", "item_name": "CV 전선", "unit": "M"}],
        },
    ).json()
    tpl_non_ea_id = c_tpl["id"]

    # 2. Upload PDF drawing
    pdf_bytes = Path("fixtures/pdf/pdf_b_2page_electrical.pdf").read_bytes()
    pdf_up = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("pdf_b_2page_electrical.pdf", pdf_bytes, "application/pdf")},
    )
    assert pdf_up.status_code == 200
    pdf_dwg_id = pdf_up.json()["drawing_id"]

    # Set initial valid EA mapping on PDF drawing
    client.put(
        f"/api/projects/{proj_id}/drawings/{pdf_dwg_id}/mappings",
        json=[{"block_name": "T1", "item_code": "SWITCH", "unit": "EA"}],
    )

    # 2a. Reject direct PUT non-EA mapping on PDF
    bad_put_pdf = client.put(
        f"/api/projects/{proj_id}/drawings/{pdf_dwg_id}/mappings",
        json=[{"block_name": "T1", "item_code": "CABLE", "unit": "M"}],
    )
    assert bad_put_pdf.status_code == 400
    assert "PDF 도면" in bad_put_pdf.json()["detail"]
    assert "EA" in bad_put_pdf.json()["detail"]

    # 2b. Reject apply-template non-EA on PDF (P0-MUNIT-03)
    bad_tpl_pdf = client.post(
        f"/api/projects/{proj_id}/drawings/{pdf_dwg_id}/apply-template",
        json={"template_id": tpl_non_ea_id, "merge_strategy": "MERGE"},
    )
    assert bad_tpl_pdf.status_code == 400
    assert "PDF 도면" in bad_tpl_pdf.json()["detail"]

    # Fail-closed check: verify original T1 mapping is unchanged
    pdf_maps = client.get(f"/api/projects/{proj_id}/drawings/{pdf_dwg_id}/mappings").json()
    assert len(pdf_maps) == 1
    assert pdf_maps[0]["block_name"] == "T1"
    assert pdf_maps[0]["unit"] == "EA"

    # 3. Upload IMAGE drawing
    img_bytes = Path("fixtures/images/floorplan_sample.png").read_bytes()
    img_up = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("floorplan_sample.png", img_bytes, "image/png")},
    )
    assert img_up.status_code == 200
    img_dwg_id = img_up.json()["drawing_id"]

    # Set initial valid EA mapping on Image drawing
    client.put(
        f"/api/projects/{proj_id}/drawings/{img_dwg_id}/mappings",
        json=[{"block_name": "IMG_T1", "item_code": "LIGHT", "unit": "EA"}],
    )

    # 3a. Reject direct PUT non-EA mapping on IMAGE
    bad_put_img = client.put(
        f"/api/projects/{proj_id}/drawings/{img_dwg_id}/mappings",
        json=[{"block_name": "IMG_T1", "item_code": "CONDUIT", "unit": "M"}],
    )
    assert bad_put_img.status_code == 400
    assert "이미지 도면" in bad_put_img.json()["detail"]
    assert "EA" in bad_put_img.json()["detail"]

    # 3b. Reject apply-template non-EA on IMAGE (P0-MUNIT-03)
    bad_tpl_img = client.post(
        f"/api/projects/{proj_id}/drawings/{img_dwg_id}/apply-template",
        json={"template_id": tpl_non_ea_id, "merge_strategy": "MERGE"},
    )
    assert bad_tpl_img.status_code == 400
    assert "이미지 도면" in bad_tpl_img.json()["detail"]

    # Fail-closed check: verify original IMG_T1 mapping is unchanged
    img_maps = client.get(f"/api/projects/{proj_id}/drawings/{img_dwg_id}/mappings").json()
    assert len(img_maps) == 1
    assert img_maps[0]["block_name"] == "IMG_T1"
    assert img_maps[0]["unit"] == "EA"

    # 4. Valid EA template succeeds on both PDF and IMAGE
    ok_pdf = client.post(
        f"/api/projects/{proj_id}/drawings/{pdf_dwg_id}/apply-template",
        json={"template_id": "tpl_standard_elec", "merge_strategy": "MERGE"},
    )
    assert ok_pdf.status_code == 200

    ok_img = client.post(
        f"/api/projects/{proj_id}/drawings/{img_dwg_id}/apply-template",
        json={"template_id": "tpl_standard_elec", "merge_strategy": "MERGE"},
    )
    assert ok_img.status_code == 200


def test_seed_demo_project_idempotency_and_schema(client: TestClient) -> None:
    """Verify demo project seed response schema and safe idempotency."""
    resp1 = client.post("/api/demo/seed")
    assert resp1.status_code == 200
    data1 = resp1.json()
    assert "project_id" in data1
    assert "project_name" in data1
    assert "drawings" in data1
    for d in data1["drawings"]:
        assert "drawing_id" in d
        assert "source_type" in d
        assert "filename" in d
        assert d["status"] == "ready"

    # Second call should reuse the project cleanly without failing
    resp2 = client.post("/api/demo/seed")
    assert resp2.status_code == 200
    data2 = resp2.json()
    assert data2["project_id"] == data1["project_id"]
    assert len(data2["drawings"]) == 3
