"""Integration tests for the FastAPI service and Takeoff workflows."""

from __future__ import annotations

import csv
import hashlib
import io
import json
from pathlib import Path

import openpyxl
import pytest
from eai_api.app import app
from fastapi.testclient import TestClient

FIXTURES_DIR = Path(__file__).resolve().parent.parent.parent.parent / "fixtures" / "dxf"


@pytest.fixture
def client(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> TestClient:
    monkeypatch.setenv("EAI_DATA_DIR", str(tmp_path / "data"))
    # Re-init storage with the temporary path
    from eai_api import app as app_module
    from eai_api.storage import Storage

    app_module.storage = Storage(root=tmp_path / "data")
    return TestClient(app)


def test_health(client: TestClient) -> None:
    res = client.get("/api/health")
    assert res.status_code == 200
    assert res.json() == {"status": "ok", "version": "0.1.0"}


def test_parse_direct_and_errors(client: TestClient) -> None:
    # 1. Reject fake DWG
    fake_dwg = FIXTURES_DIR / "fake_dwg.dxf"
    res = client.post(
        "/api/dxf/parse-direct",
        files={"file": (fake_dwg.name, fake_dwg.read_bytes(), "application/octet-stream")},
    )
    assert res.status_code == 400
    assert res.json()["code"] == "DWG_NOT_SUPPORTED"

    # 2. Reject non-DXF
    not_dxf = FIXTURES_DIR / "not_a_dxf.dxf"
    res = client.post(
        "/api/dxf/parse-direct",
        files={"file": (not_dxf.name, not_dxf.read_bytes(), "text/plain")},
    )
    assert res.status_code == 400
    assert res.json()["code"] == "NOT_A_DXF"

    # 3. Direct parse valid basic_devices
    valid = FIXTURES_DIR / "basic_devices.dxf"
    res = client.post(
        "/api/dxf/parse-direct",
        files={"file": (valid.name, valid.read_bytes(), "application/octet-stream")},
    )
    assert res.status_code == 200
    data = res.json()
    assert data["source"]["filename"] == "basic_devices.dxf"
    assert len(data["objects"]) == 18


def test_full_takeoff_lifecycle(client: TestClient) -> None:
    # 1. Create project
    proj_res = client.post("/api/projects", json={"name": "테스트 전기공사 프로젝트"})
    assert proj_res.status_code == 200
    project_id = proj_res.json()["id"]

    # 2. Upload drawing
    dxf_file = FIXTURES_DIR / "basic_devices.dxf"
    upload_res = client.post(
        f"/api/projects/{project_id}/drawings",
        files={"file": (dxf_file.name, dxf_file.read_bytes(), "application/octet-stream")},
    )
    assert upload_res.status_code == 200
    drawing_id = upload_res.json()["drawing_id"]

    # 3. Get normalized drawing
    get_dxf_res = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}")
    assert get_dxf_res.status_code == 200
    drawing_data = get_dxf_res.json()
    assert len(drawing_data["objects"]) == 18

    # 4. Set mappings
    mappings = [
        {"block_name": "LT-DL", "item_code": "LIGHT", "item_name": "다운라이트 LED 6인치", "unit": "EA"},
        {"block_name": "OT-2P", "item_code": "OUTLET", "item_name": "매입 2구 콘센트", "unit": "EA"},
        {"block_name": "SW-1", "item_code": "SWITCH", "item_name": "1로 텀블러 스위치", "unit": "EA"},
    ]
    map_res = client.put(f"/api/projects/{project_id}/drawings/{drawing_id}/mappings", json=mappings)
    assert map_res.status_code == 200

    # 5. Check Takeoff results: LIGHT=6, OUTLET=4, SWITCH=3
    takeoff_res = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff")
    assert takeoff_res.status_code == 200
    takeoff_data = takeoff_res.json()
    sum_map = {row["item_code"]: row["quantity"] for row in takeoff_data["summary"]}
    assert sum_map == {"LIGHT": 6, "OUTLET": 4, "SWITCH": 3}

    # 6. User Review Override: exclude 1 LIGHT object
    light_evidence = [e for e in takeoff_data["evidence"] if e["mapped_item_code"] == "LIGHT"]
    target_obj_id = light_evidence[0]["source_object_id"]

    client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/reviews",
        json=[
            {
                "source_object_id": target_obj_id,
                "inclusion_status": "EXCLUDED",
                "review_status": "REVIEWED",
                "reason": "오탐 제외 (도면 표기 불일치)",
            }
        ],
    )

    # Check updated Takeoff: LIGHT is now 5
    takeoff_res2 = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff")
    sum_map2 = {row["item_code"]: row["quantity"] for row in takeoff_res2.json()["summary"]}
    assert sum_map2["LIGHT"] == 5

    # 7. Add Manual Object: 1 missing LIGHT
    client.post(
        f"/api/projects/{project_id}/drawings/{drawing_id}/manual-objects",
        json={
            "id": "manual_lt_1",
            "drawing_id": drawing_id,
            "x": 2500.0,
            "y": 2500.0,
            "item_code": "LIGHT",
            "item_name": "다운라이트 LED 6인치",
            "note": "누락 기구 현장 추가",
        },
    )

    # Check updated Takeoff: LIGHT is back to 6, method is MIXED
    takeoff_res3 = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff")
    light_row = next(r for r in takeoff_res3.json()["summary"] if r["item_code"] == "LIGHT")
    assert light_row["quantity"] == 6
    assert light_row["quantity_method"] == "MIXED"

    # 8. Export CSV summary & evidence
    csv_sum = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/export/csv?type=summary")
    assert csv_sum.status_code == 200
    assert "item_code,item_name" in csv_sum.text
    assert "LIGHT" in csv_sum.text

    csv_ev = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/export/csv?type=evidence")
    assert csv_ev.status_code == 200
    assert "source_object_id,drawing_id" in csv_ev.text
    assert target_obj_id in csv_ev.text

    # 9. Export XLSX
    xlsx_res = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/export/xlsx")
    assert xlsx_res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_res.content))
    assert "Takeoff Summary" in wb.sheetnames
    assert "Takeoff Evidence" in wb.sheetnames

    # Check that summary quantity in Excel matches
    ws_sum = wb["Takeoff Summary"]
    found_light = False
    for row in ws_sum.iter_rows(values_only=True):
        if row[0] == "LIGHT":
            assert row[4] == 6  # quantity
            found_light = True
    assert found_light


def test_list_drawings_and_reopen(client: TestClient) -> None:
    # 1. Create project
    proj_res = client.post("/api/projects", json={"name": "도면 목록 테스트"})
    assert proj_res.status_code == 200
    project_id = proj_res.json()["id"]

    # 2. Upload first drawing
    dxf1 = FIXTURES_DIR / "basic_devices.dxf"
    res1 = client.post(
        f"/api/projects/{project_id}/drawings?revision=REV0",
        files={"file": (dxf1.name, dxf1.read_bytes(), "application/octet-stream")},
    )
    assert res1.status_code == 200
    d1_id = res1.json()["drawing_id"]

    # 3. Upload second drawing
    dxf2 = FIXTURES_DIR / "hidden_layers.dxf"
    res2 = client.post(
        f"/api/projects/{project_id}/drawings?revision=REV1",
        files={"file": (dxf2.name, dxf2.read_bytes(), "application/octet-stream")},
    )
    assert res2.status_code == 200
    d2_id = res2.json()["drawing_id"]

    # 4. List drawings
    list_res = client.get(f"/api/projects/{project_id}/drawings")
    assert list_res.status_code == 200
    drawings = list_res.json()
    assert len(drawings) == 2
    drawing_ids = [d["id"] for d in drawings]
    assert d1_id in drawing_ids
    assert d2_id in drawing_ids

    # 5. Reopen saved drawing
    reopened = client.get(f"/api/projects/{project_id}/drawings/{d1_id}")
    assert reopened.status_code == 200
    data = reopened.json()
    assert data["drawing_id"] == d1_id
    assert len(data["objects"]) == 18


def test_structural_blocking_defense_ac_06(client: TestClient) -> None:
    # AC-06: Structurally invalid objects (XREF, missing block definitions, circular)
    # MUST NEVER be forcefully included into CAD counts even if user sets INCLUDED.
    proj_res = client.post("/api/projects", json={"name": "XREF 방어 테스트"})
    project_id = proj_res.json()["id"]

    xref_file = FIXTURES_DIR / "xref_reference.dxf"
    upload_res = client.post(
        f"/api/projects/{project_id}/drawings",
        files={"file": (xref_file.name, xref_file.read_bytes(), "application/octet-stream")},
    )
    assert upload_res.status_code == 200
    drawing_id = upload_res.json()["drawing_id"]

    # Map both LT-DL and ARCH_BASE (which is an XREF) to LIGHT
    client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/mappings",
        json=[
            {"block_name": "LT-DL", "item_code": "LIGHT", "unit": "EA"},
            {"block_name": "ARCH_BASE", "item_code": "LIGHT", "unit": "EA"},
        ],
    )

    # Initial takeoff: ARCH_BASE is blocked, so LIGHT is only 1 (from LT-DL)
    t1 = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    sum1 = {row["item_code"]: row["quantity"] for row in t1["summary"]}
    assert sum1["LIGHT"] == 1

    # Find ARCH_BASE object
    xref_ev = next(e for e in t1["evidence"] if e["block_name"] == "ARCH_BASE")
    assert xref_ev["inclusion_status"] == "EXCLUDED"
    assert xref_ev["review_status"] == "UNRESOLVED"

    # Attempt to forcefully override with INCLUDED
    client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/reviews",
        json=[
            {
                "source_object_id": xref_ev["source_object_id"],
                "inclusion_status": "INCLUDED",
                "review_status": "REVIEWED",
                "mapped_item_code": "LIGHT",
                "reason": "강제 포함 시도",
            }
        ],
    )

    # Re-check takeoff: LIGHT must STILL be 1, defense active!
    t2 = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    sum2 = {row["item_code"]: row["quantity"] for row in t2["summary"]}
    assert sum2["LIGHT"] == 1

    xref_ev2 = next(e for e in t2["evidence"] if e["source_object_id"] == xref_ev["source_object_id"])
    assert xref_ev2["inclusion_status"] == "EXCLUDED"
    assert xref_ev2["review_status"] == "UNRESOLVED"
    assert "구조적 결함" in xref_ev2["review_reason"]

    # AC-06/FR-011: Verify summary row explicitly marks LIGHT as UNRESOLVED due to unresolved XREF
    light_sum = next(r for r in t2["summary"] if r["item_code"] == "LIGHT")
    assert light_sum["review_status"] == "UNRESOLVED"
    assert "미해결 객체" in light_sum["notes"]


def test_binary_dxf_direct_parse_and_upload(client: TestClient) -> None:
    # 1. Direct parse real AutoCAD binary DXF
    bin_file = FIXTURES_DIR / "binary_devices.dxf"
    res = client.post(
        "/api/dxf/parse-direct",
        files={"file": (bin_file.name, bin_file.read_bytes(), "application/octet-stream")},
    )
    assert res.status_code == 200
    data = res.json()
    assert len(data["objects"]) == 18
    assert data["source"]["filename"] == "binary_devices.dxf"

    # 2. Upload to project and calculate takeoff
    proj_res = client.post("/api/projects", json={"name": "바이너리 DXF 프로젝트"})
    project_id = proj_res.json()["id"]

    upload_res = client.post(
        f"/api/projects/{project_id}/drawings",
        files={"file": (bin_file.name, bin_file.read_bytes(), "application/octet-stream")},
    )
    assert upload_res.status_code == 200
    drawing_id = upload_res.json()["drawing_id"]

    # Map blocks
    client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/mappings",
        json=[
            {"block_name": "LT-DL", "item_code": "LIGHT", "unit": "EA"},
            {"block_name": "OT-2P", "item_code": "OUTLET", "unit": "EA"},
            {"block_name": "SW-1", "item_code": "SWITCH", "unit": "EA"},
        ],
    )

    takeoff = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    sum_map = {row["item_code"]: row["quantity"] for row in takeoff["summary"]}
    assert sum_map == {"LIGHT": 6, "OUTLET": 4, "SWITCH": 3}


def test_manual_object_override_and_delete(client: TestClient) -> None:
    proj_res = client.post("/api/projects", json={"name": "수동 기구 관리 테스트"})
    project_id = proj_res.json()["id"]

    dxf_file = FIXTURES_DIR / "basic_devices.dxf"
    upload_res = client.post(
        f"/api/projects/{project_id}/drawings",
        files={"file": (dxf_file.name, dxf_file.read_bytes(), "application/octet-stream")},
    )
    drawing_id = upload_res.json()["drawing_id"]

    # Map LT-DL to LIGHT
    client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/mappings",
        json=[{"block_name": "LT-DL", "item_code": "LIGHT", "unit": "EA"}],
    )

    # 1. Add manual object as LIGHT
    manual_id = "manual_test_100"
    client.post(
        f"/api/projects/{project_id}/drawings/{drawing_id}/manual-objects",
        json={
            "id": manual_id,
            "drawing_id": drawing_id,
            "x": 100.0,
            "y": 200.0,
            "item_code": "LIGHT",
            "item_name": "수동 추가 조명",
            "note": "현장 추가",
        },
    )

    t1 = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    assert next(r for r in t1["summary"] if r["item_code"] == "LIGHT")["quantity"] == 7

    # 2. Apply review override to manual object (change item_code to OUTLET)
    client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/reviews",
        json=[
            {
                "source_object_id": manual_id,
                "inclusion_status": "MANUAL_ADDED",
                "review_status": "REVIEWED",
                "mapped_item_code": "OUTLET",
                "reason": "조명이 아닌 콘센트로 정정",
            }
        ],
    )

    t2 = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    assert next(r for r in t2["summary"] if r["item_code"] == "LIGHT")["quantity"] == 6
    assert next(r for r in t2["summary"] if r["item_code"] == "OUTLET")["quantity"] == 1

    # 3. Delete manual object
    del_res = client.delete(f"/api/projects/{project_id}/drawings/{drawing_id}/manual-objects/{manual_id}")
    assert del_res.status_code == 200

    t3 = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    assert next(r for r in t3["summary"] if r["item_code"] == "LIGHT")["quantity"] == 6
    outlet_row = next((r for r in t3["summary"] if r["item_code"] == "OUTLET"), None)
    assert outlet_row is None or outlet_row["quantity"] == 0


def test_review_confirm_preserves_included_quantity(client: TestClient) -> None:
    # Verifies the "검수 확인" workflow: marking INCLUDED as REVIEWED preserves quantity
    proj_res = client.post("/api/projects", json={"name": "검수 확인 테스트"})
    project_id = proj_res.json()["id"]

    dxf_file = FIXTURES_DIR / "basic_devices.dxf"
    upload_res = client.post(
        f"/api/projects/{project_id}/drawings",
        files={"file": (dxf_file.name, dxf_file.read_bytes(), "application/octet-stream")},
    )
    drawing_id = upload_res.json()["drawing_id"]

    client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/mappings",
        json=[{"block_name": "LT-DL", "item_code": "LIGHT", "unit": "EA"}],
    )

    t1 = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    assert next(r for r in t1["summary"] if r["item_code"] == "LIGHT")["quantity"] == 6

    # Select first LIGHT evidence row
    first_light = next(e for e in t1["evidence"] if e["mapped_item_code"] == "LIGHT")
    assert first_light["review_status"] == "UNREVIEWED"

    # Confirm review
    client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/reviews",
        json=[
            {
                "source_object_id": first_light["source_object_id"],
                "inclusion_status": "INCLUDED",
                "review_status": "REVIEWED",
                "reason": "현장 확인 완료",
            }
        ],
    )

    t2 = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    assert next(r for r in t2["summary"] if r["item_code"] == "LIGHT")["quantity"] == 6
    updated_ev = next(e for e in t2["evidence"] if e["source_object_id"] == first_light["source_object_id"])
    assert updated_ev["review_status"] == "REVIEWED"
    assert updated_ev["inclusion_status"] == "INCLUDED"
    assert updated_ev["review_reason"] == "현장 확인 완료"


def test_export_formula_injection_defense(client: TestClient) -> None:
    # Verifies Issue #7 defense: OWASP CSV & Excel formula injection
    proj_res = client.post("/api/projects", json={"name": "수식 인젝션 방어 테스트"})
    project_id = proj_res.json()["id"]

    dxf_file = FIXTURES_DIR / "basic_devices.dxf"
    upload_res = client.post(
        f"/api/projects/{project_id}/drawings",
        files={"file": (dxf_file.name, dxf_file.read_bytes(), "application/octet-stream")},
    )
    drawing_id = upload_res.json()["drawing_id"]

    # Inject formula characters (=, +, -, @) in mappings and review notes
    client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/mappings",
        json=[
            {
                "block_name": "LT-DL",
                "item_code": "=cmd|' /C calc'!A0",
                "item_name": "@SUM(1,1)",
                "unit": "EA",
            }
        ],
    )

    t1 = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    first_ev = t1["evidence"][0]

    client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/reviews",
        json=[
            {
                "source_object_id": first_ev["source_object_id"],
                "inclusion_status": "INCLUDED",
                "review_status": "REVIEWED",
                "reason": "+cmd|' /C calc'!A0",
            }
        ],
    )

    # 1. Check CSV Summary parsed line-by-line via csv.reader
    csv_sum_res = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/export/csv?type=summary")
    assert csv_sum_res.status_code == 200
    reader_sum = list(csv.reader(io.StringIO(csv_sum_res.content.decode("utf-8-sig"))))
    assert len(reader_sum) >= 2
    header_sum = reader_sum[0]
    assert header_sum == [
        "item_code",
        "item_name",
        "specification",
        "unit",
        "quantity",
        "quantity_method",
        "review_status",
        "drawing_id",
        "drawing_name",
        "drawing_revision",
        "exported_at",
        "completeness_status",
        "notes",
    ]
    # Check row 1 fields
    row_sum = reader_sum[1]
    assert row_sum[0].startswith("'=")
    assert row_sum[1].startswith("'@")
    assert row_sum[7] == drawing_id
    assert row_sum[9] == "REV0"

    # 2. Check CSV Evidence parsed line-by-line via csv.reader
    csv_ev_res = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/export/csv?type=evidence")
    assert csv_ev_res.status_code == 200
    reader_ev = list(csv.reader(io.StringIO(csv_ev_res.content.decode("utf-8-sig"))))
    assert len(reader_ev) >= 2
    header_ev = reader_ev[0]
    assert header_ev == [
        "source_object_id",
        "drawing_id",
        "parent_source_id",
        "block_name",
        "layer_name",
        "entity_type",
        "x",
        "y",
        "mapped_item_code",
        "inclusion_status",
        "review_status",
        "review_reason",
        "drawing_revision",
        "exported_at",
    ]
    # First row reason must have leading quote escape
    row_ev = reader_ev[1]
    assert row_ev[11].startswith("'+")

    # 3. Check Excel XLSX deep structure and cell.data_type != 'f'
    xlsx_res = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/export/xlsx")
    assert xlsx_res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_res.content))
    assert "Takeoff Summary" in wb.sheetnames
    assert "Takeoff Evidence" in wb.sheetnames

    for sheet_name in ("Takeoff Summary", "Takeoff Evidence"):
        ws = wb[sheet_name]
        for row_cells in ws.iter_rows():
            for c in row_cells:
                assert c.data_type != "f", f"Formula cell detected in {sheet_name} at {c.coordinate}: {c.value}"

    ws_sum = wb["Takeoff Summary"]
    assert str(ws_sum.cell(row=2, column=1).value).startswith("'=")
    assert str(ws_sum.cell(row=2, column=2).value).startswith("'@")


def test_raw_dxf_preservation_and_download(client: TestClient) -> None:
    """Issue #5 / QA-DEF-01: Raw DXF bytes must be preserved and retrievable byte-for-byte."""
    proj = client.post("/api/projects", json={"name": "원본 보존 검증 프로젝트"}).json()
    project_id = proj["id"]

    dxf_file = FIXTURES_DIR / "basic_devices.dxf"
    original_bytes = dxf_file.read_bytes()
    expected_hash = hashlib.sha256(original_bytes).hexdigest()

    upload_res = client.post(
        f"/api/projects/{project_id}/drawings",
        files={"file": (dxf_file.name, original_bytes, "application/octet-stream")},
    )
    assert upload_res.status_code == 200
    drawing_info = upload_res.json()
    assert drawing_info["raw_saved"] is True
    assert drawing_info["file_size_bytes"] == len(original_bytes)
    drawing_id = drawing_info["drawing_id"]

    # Download raw drawing via API
    raw_res = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/raw")
    assert raw_res.status_code == 200
    assert raw_res.content == original_bytes
    assert hashlib.sha256(raw_res.content).hexdigest() == expected_hash
    assert "attachment" in raw_res.headers.get("content-disposition", "")


def test_streaming_upload_boundary_and_413(client: TestClient) -> None:
    """Issue #5 / QA-DEF-02: Bounded streaming upload must reject files > 50 MiB with HTTP 413."""
    proj = client.post("/api/projects", json={"name": "용량 제한 검증"}).json()
    project_id = proj["id"]

    # 1. 0 bytes -> returns HTTP 400 (empty file rejected by loader)
    zero_res = client.post(
        f"/api/projects/{project_id}/drawings",
        files={"file": ("empty.dxf", b"", "application/octet-stream")},
    )
    assert zero_res.status_code == 400
    assert zero_res.json()["code"] == "EMPTY_FILE"

    # 2. 50 MiB + 1 byte -> returns HTTP 413 Payload Too Large
    oversized_len = 50 * 1024 * 1024 + 1
    # Use generator-like stream of dummy bytes to simulate oversized stream
    oversized_data = b"0\nSECTION\n" + b" " * (oversized_len - 10)
    over_res = client.post(
        f"/api/projects/{project_id}/drawings",
        files={"file": ("huge.dxf", oversized_data, "application/octet-stream")},
    )
    assert over_res.status_code == 413
    assert "50 MiB" in over_res.json()["detail"]


def test_provisional_and_unresolved_contracts(client: TestClient) -> None:
    """Issue #9 / DEC-010~014: Provisional quantity and unresolved items contract."""
    proj = client.post("/api/projects", json={"name": "잠정 산출 검증"}).json()
    project_id = proj["id"]

    dxf_file = FIXTURES_DIR / "basic_devices.dxf"
    upload_res = client.post(
        f"/api/projects/{project_id}/drawings",
        files={"file": (dxf_file.name, dxf_file.read_bytes(), "application/octet-stream")},
    )
    drawing_id = upload_res.json()["drawing_id"]

    # No mappings yet: all items are UNMAPPED -> must be provisional with unresolved_count > 0
    t0 = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    assert t0["is_provisional"] is True
    assert t0["unresolved_count"] > 0

    # Map LT-DL to LIGHT, but others unmapped
    client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/mappings",
        json=[{"block_name": "LT-DL", "item_code": "LIGHT", "item_name": "다운라이트", "unit": "EA"}],
    )
    t1 = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    # Still provisional because other blocks (OT-2P, SW-1, etc.) remain unmapped
    assert t1["is_provisional"] is True
    assert t1["unresolved_count"] > 0


def test_dynamic_and_anonymous_block_override_defense_qa_09_c(client: TestClient) -> None:
    """Issue #9 / QA-09-C: Dynamic/anonymous blocks must NEVER be forcefully included via user review override."""
    proj = client.post("/api/projects", json={"name": "동적 블록 오버라이드 방어"}).json()
    project_id = proj["id"]

    dxf_file = FIXTURES_DIR / "anonymous_and_dynamic.dxf"
    upload_res = client.post(
        f"/api/projects/{project_id}/drawings",
        files={"file": (dxf_file.name, dxf_file.read_bytes(), "application/octet-stream")},
    )
    drawing_id = upload_res.json()["drawing_id"]

    # Map LT-DYN (which has DYNAMIC_BLOCK warning)
    client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/mappings",
        json=[{"block_name": "LT-DYN", "item_code": "LIGHT", "unit": "EA"}],
    )

    t0 = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    sum0 = {r["item_code"]: r["quantity"] for r in t0["summary"]}
    assert sum0.get("LIGHT", 0) == 0

    dyn_ev = next(e for e in t0["evidence"] if e["block_name"] == "*U2" or "LT-DYN" in e["mapped_item_code"] or e["mapped_item_code"] == "LIGHT")
    assert dyn_ev["inclusion_status"] == "EXCLUDED"
    assert dyn_ev["review_status"] == "UNRESOLVED"

    # ATTACK: User attempts to force override with INCLUDED + REVIEWED
    client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/reviews",
        json=[
            {
                "source_object_id": dyn_ev["source_object_id"],
                "inclusion_status": "INCLUDED",
                "review_status": "REVIEWED",
                "mapped_item_code": "LIGHT",
                "reason": "동적 블록 강제 포함 시도",
            }
        ],
    )

    t1 = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    sum1 = {r["item_code"]: r["quantity"] for r in t1["summary"]}
    assert sum1.get("LIGHT", 0) == 0  # Must remain 0 EA!
    dyn_ev1 = next(e for e in t1["evidence"] if e["source_object_id"] == dyn_ev["source_object_id"])
    assert dyn_ev1["inclusion_status"] == "EXCLUDED"
    assert dyn_ev1["review_status"] == "UNRESOLVED"
    assert "동적/익명 블록" in dyn_ev1["review_reason"]
    assert t1["unresolved_count"] > 0
    assert t1["is_provisional"] is True


def test_nested_blocks_withholding_and_manual_addition_qa_09_d(client: TestClient) -> None:
    """Issue #9 / QA-09-D: In nested_blocks.dxf, auto count is strictly LIGHT=1; overrides are blocked; manual additions add conservatively."""
    proj = client.post("/api/projects", json={"name": "중첩 계층 보류 및 수동 산출 검증"}).json()
    project_id = proj["id"]

    nested_file = FIXTURES_DIR / "nested_blocks.dxf"
    upload_res = client.post(
        f"/api/projects/{project_id}/drawings",
        files={"file": (nested_file.name, nested_file.read_bytes(), "application/octet-stream")},
    )
    drawing_id = upload_res.json()["drawing_id"]

    # 1. Map both parent assembly LT-GROUP and child fixture LT-DL to LIGHT
    client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/mappings",
        json=[
            {"block_name": "LT-GROUP", "item_code": "LIGHT", "unit": "EA"},
            {"block_name": "LT-DL", "item_code": "LIGHT", "unit": "EA"},
        ],
    )

    t0 = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    sum0 = {r["item_code"]: r["quantity"] for r in t0["summary"]}
    # Only 1 independent top-level LT-DL is counted! Parent assemblies (2) & nested children (4) are withheld.
    assert sum0.get("LIGHT") == 1

    # 2. Find parent object and child object
    nested_child = next(e for e in t0["evidence"] if e["block_name"] == "LT-DL" and e["parent_source_id"] is not None)
    parent_obj = next(e for e in t0["evidence"] if e["block_name"] == "LT-GROUP")

    # 3. ATTACK: User attempts review override to force INCLUDED on parent and child
    client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/reviews",
        json=[
            {
                "source_object_id": parent_obj["source_object_id"],
                "inclusion_status": "INCLUDED",
                "review_status": "REVIEWED",
                "mapped_item_code": "LIGHT",
                "reason": "상위 어셈블리 강제 포함 시도",
            },
            {
                "source_object_id": nested_child["source_object_id"],
                "inclusion_status": "INCLUDED",
                "review_status": "REVIEWED",
                "mapped_item_code": "LIGHT",
                "reason": "하위 중첩 블록 강제 포함 시도",
            },
        ],
    )

    t1 = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    sum1 = {r["item_code"]: r["quantity"] for r in t1["summary"]}
    # CAD automatic count remains STRICTLY 1 EA!
    assert sum1.get("LIGHT") == 1

    parent_ev1 = next(e for e in t1["evidence"] if e["source_object_id"] == parent_obj["source_object_id"])
    child_ev1 = next(e for e in t1["evidence"] if e["source_object_id"] == nested_child["source_object_id"])
    assert parent_ev1["inclusion_status"] == "EXCLUDED"
    assert parent_ev1["review_status"] == "UNRESOLVED"
    assert "중첩 블록 계층" in parent_ev1["review_reason"]

    assert child_ev1["inclusion_status"] == "EXCLUDED"
    assert child_ev1["review_status"] == "UNRESOLVED"
    assert "중첩 블록 계층" in child_ev1["review_reason"]

    # 4. FIELD VERIFICATION: Engineer verifies 4 actual lamps in field and adds 4 manual markers
    for i in range(4):
        mo_res = client.post(
            f"/api/projects/{project_id}/drawings/{drawing_id}/manual-objects",
            json={
                "id": f"manual_lt_{i + 1}",
                "drawing_id": drawing_id,
                "item_code": "LIGHT",
                "item_name": "다운라이트 LED 6인치",
                "x": 100.0 * (i + 1),
                "y": 200.0 * (i + 1),
                "note": f"현장 검수 확인 조명 #{i + 1}",
            },
        )
        assert mo_res.status_code == 200

    t2 = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    sum2 = {r["item_code"]: r["quantity"] for r in t2["summary"]}
    # Total provisional LIGHT quantity is now 1 (CAD) + 4 (MANUAL) = 5 EA!
    assert sum2.get("LIGHT") == 5

    light_summary_row = next(r for r in t2["summary"] if r["item_code"] == "LIGHT")
    assert light_summary_row["quantity_method"] == "MIXED"

    # Verify ID uniqueness and evidence integrity
    evidence_ids = [e["source_object_id"] for e in t2["evidence"]]
    assert len(evidence_ids) == len(set(evidence_ids)), "All evidence IDs must be strictly unique"

    manual_evs = [e for e in t2["evidence"] if e["inclusion_status"] == "MANUAL_ADDED"]
    assert len(manual_evs) == 4
    assert t2["is_provisional"] is True
    assert t2["unresolved_count"] >= 6


def test_deep_openpyxl_and_csv_reader_verification(client: TestClient) -> None:
    """Issue #7: Thorough openpyxl workbook parsing and standard RFC csv.reader verification."""
    proj = client.post("/api/projects", json={"name": "내보내기 신뢰성 검증"}).json()
    project_id = proj["id"]

    dxf_file = FIXTURES_DIR / "basic_devices.dxf"
    upload_res = client.post(
        f"/api/projects/{project_id}/drawings",
        files={"file": (dxf_file.name, dxf_file.read_bytes(), "application/octet-stream")},
    )
    drawing_id = upload_res.json()["drawing_id"]

    # Map LT-DL, OT-2P, SW-1
    client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/mappings",
        json=[
            {"block_name": "LT-DL", "item_code": "LIGHT", "item_name": "다운라이트", "unit": "EA"},
            {"block_name": "OT-2P", "item_code": "OUTLET", "item_name": "매입콘센트 2구", "unit": "EA"},
            {"block_name": "SW-1", "item_code": "SWITCH", "item_name": "1로 텀블러 스위치", "unit": "EA"},
        ],
    )

    # 1. Verify CSV Summary with standard python csv.reader
    csv_sum_bytes = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/export/csv?type=summary").content
    csv_text = csv_sum_bytes.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(csv_text))
    rows = list(reader)

    header = rows[0]
    assert header[0] == "item_code"
    assert header[4] == "quantity"
    assert header[10] == "exported_at"
    assert header[11] == "completeness_status"
    assert header[12] == "notes"

    data_rows = {r[0]: r for r in rows[1:]}
    assert int(data_rows["LIGHT"][4]) == 6
    assert int(data_rows["OUTLET"][4]) == 4
    assert int(data_rows["SWITCH"][4]) == 3
    assert "잠정" in data_rows["LIGHT"][11]

    # 2. Verify Excel with openpyxl.load_workbook
    xlsx_bytes = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/export/xlsx").content
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert wb.sheetnames == ["Takeoff Summary", "Takeoff Evidence"]

    # Inspect Takeoff Summary sheet
    ws_sum = wb["Takeoff Summary"]
    assert ws_sum.cell(row=1, column=1).value == "품목코드 (item_code)"
    assert ws_sum.cell(row=1, column=5).value == "산출 수량 (quantity)"
    assert ws_sum.cell(row=1, column=12).value == "산출 상태 (completeness_status)"

    summary_data = {}
    for r in range(2, ws_sum.max_row + 1):
        item_code = ws_sum.cell(row=r, column=1).value
        qty = ws_sum.cell(row=r, column=5).value
        status = ws_sum.cell(row=r, column=12).value
        summary_data[item_code] = {"qty": qty, "status": status}
        # Invariant: no formula cells
        for c in range(1, ws_sum.max_column + 1):
            assert ws_sum.cell(row=r, column=c).data_type != "f"

    assert summary_data["LIGHT"]["qty"] == 6
    assert summary_data["OUTLET"]["qty"] == 4
    assert summary_data["SWITCH"]["qty"] == 3
    assert "잠정" in summary_data["LIGHT"]["status"]

    # Inspect Takeoff Evidence sheet
    ws_ev = wb["Takeoff Evidence"]
    assert ws_ev.max_row == 14  # 1 header + 13 objects
    for r in range(1, ws_ev.max_row + 1):
        for c in range(1, ws_ev.max_column + 1):
            cell = ws_ev.cell(row=r, column=c)
            assert cell.data_type != "f"


def test_conservative_provisional_policy_without_scope_verification(client: TestClient) -> None:
    """Issue #9 item 4: Even if all items are REVIEWED, is_provisional remains True in P0 without scope verification."""
    proj = client.post("/api/projects", json={"name": "보수적 잠정 정책"}).json()
    project_id = proj["id"]

    dxf_file = FIXTURES_DIR / "basic_devices.dxf"
    upload_res = client.post(
        f"/api/projects/{project_id}/drawings",
        files={"file": (dxf_file.name, dxf_file.read_bytes(), "application/octet-stream")},
    )
    drawing_id = upload_res.json()["drawing_id"]

    client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/mappings",
        json=[
            {"block_name": "LT-DL", "item_code": "LIGHT"},
            {"block_name": "OT-2P", "item_code": "OUTLET"},
            {"block_name": "SW-1", "item_code": "SWITCH"},
        ],
    )

    t = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    # Mark all evidence as REVIEWED
    reviews_payload = [
        {
            "source_object_id": e["source_object_id"],
            "inclusion_status": e["inclusion_status"],
            "review_status": "REVIEWED",
            "mapped_item_code": e["mapped_item_code"],
            "reason": "전수 검수 완료",
        }
        for e in t["evidence"]
    ]
    client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/reviews",
        json=reviews_payload,
    )

    t_reviewed = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    assert t_reviewed["unresolved_count"] == 0
    assert all(s["review_status"] == "REVIEWED" for s in t_reviewed["summary"])
    # Without explicit count_scope_verified, is_provisional must remain True!
    assert t_reviewed["is_provisional"] is True


def test_adr_005_api_mapping_validation_and_canonical_item_contract(client: TestClient) -> None:
    """Issue #9 / ADR-005: Full API validation and fail-closed isolation for canonical item definitions."""
    proj = client.post("/api/projects", json={"name": "ADR-005 정규 품목 계약 검증"}).json()
    project_id = proj["id"]

    dxf_file = FIXTURES_DIR / "basic_devices.dxf"
    upload_res = client.post(
        f"/api/projects/{project_id}/drawings",
        files={"file": (dxf_file.name, dxf_file.read_bytes(), "application/octet-stream")},
    )
    drawing_id = upload_res.json()["drawing_id"]

    # 1. SCENARIO A: Same block, same code, conflicting units ("EA" vs "SET")
    res_a = client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/mappings",
        json=[
            {"block_name": "LT-DL", "item_code": "LIGHT", "unit": "EA", "specification": "6인치"},
            {"block_name": "LT-DL", "item_code": "LIGHT", "unit": "SET", "specification": "6인치"},
        ],
    )
    assert res_a.status_code == 400
    assert "서로 다른 단위('EA' vs 'SET')" in res_a.json()["detail"]
    assert "ADR-005" in res_a.json()["detail"]

    # 2. SCENARIO B: Same block, same code, conflicting specifications ("6인치" vs "12인치")
    res_b = client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/mappings",
        json=[
            {"block_name": "LT-DL", "item_code": "LIGHT", "unit": "EA", "specification": "6인치"},
            {"block_name": "LT-DL", "item_code": "LIGHT", "unit": "EA", "specification": "12인치"},
        ],
    )
    assert res_b.status_code == 400
    assert "서로 다른 규격('6인치' vs '12인치')" in res_b.json()["detail"]

    # 3. SCENARIO C: Different blocks, same code, conflicting units
    res_c = client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/mappings",
        json=[
            {"block_name": "LT-DL", "item_code": "LIGHT", "unit": "EA"},
            {"block_name": "OT-2P", "item_code": "LIGHT", "unit": "SET"},
        ],
    )
    assert res_c.status_code == 400
    assert "서로 다른 단위" in res_c.json()["detail"]

    # 4. SCENARIO D: Reversed order also rejected with 400
    res_d = client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/mappings",
        json=[
            {"block_name": "OT-2P", "item_code": "LIGHT", "unit": "SET"},
            {"block_name": "LT-DL", "item_code": "LIGHT", "unit": "EA"},
        ],
    )
    assert res_d.status_code == 400

    # 5. Save a valid canonical mapping: LIGHT as EA
    valid_res = client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/mappings",
        json=[
            {"block_name": "LT-DL", "item_code": "LIGHT", "unit": "EA", "specification": "6인치 LED"},
            {"block_name": "OT-2P", "item_code": "OUTLET", "unit": "EA", "specification": "매입 2구"},
        ],
    )
    assert valid_res.status_code == 200

    # 6. SCENARIO E: Manual object unit mismatch against canonical mapping fails
    mo_bad = client.post(
        f"/api/projects/{project_id}/drawings/{drawing_id}/manual-objects",
        json={
            "id": "man_err_1",
            "drawing_id": drawing_id,
            "item_code": "LIGHT",
            "unit": "SET",  # Canonical mapping is EA!
            "x": 100.0,
            "y": 100.0,
        },
    )
    assert mo_bad.status_code == 400
    assert "정규 단위('EA')와 일치하지 않습니다" in mo_bad.json()["detail"]

    # 7. SCENARIO F: Legacy / Tampered Storage Simulation (Direct file injection of conflicting mappings)
    # Bypass PUT /mappings by directly writing conflicting mappings into storage
    from eai_api.app import storage
    drawings_dir = storage.projects_dir / project_id / "drawings" / drawing_id
    tampered_mappings = [
        {"block_name": "LT-DL", "item_code": "LIGHT", "unit": "EA", "specification": "6인치"},
        {"block_name": "OT-2P", "item_code": "LIGHT", "unit": "SET", "specification": "12인치"},
    ]
    (drawings_dir / "mappings.json").write_text(json.dumps(tampered_mappings), encoding="utf-8")

    # Call GET /takeoff: must fail-closed without silent summing!
    t_tampered = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    light_row = next(r for r in t_tampered["summary"] if r["item_code"] == "LIGHT")
    assert light_row["quantity"] == 0, "Conflicted item_code MUST NOT be counted"
    assert light_row["review_status"] == "UNRESOLVED"
    assert "품목 정의 충돌" in light_row["notes"]
    assert light_row["specification"] == "[상충 규격 충돌]"
    assert light_row["unit"] == "[상충 단위 충돌]"
    assert t_tampered["is_provisional"] is True
    assert t_tampered["unresolved_count"] > 0

    # User attempting INCLUDED review override on conflicted item is blocked!
    ev_light = next(e for e in t_tampered["evidence"] if e["mapped_item_code"] == "LIGHT")
    client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/reviews",
        json=[
            {
                "source_object_id": ev_light["source_object_id"],
                "inclusion_status": "INCLUDED",
                "review_status": "REVIEWED",
                "mapped_item_code": "LIGHT",
                "reason": "충돌 품목 강제 포함 시도",
            }
        ],
    )
    t_after_override = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    ev_after = next(e for e in t_after_override["evidence"] if e["source_object_id"] == ev_light["source_object_id"])
    assert ev_after["inclusion_status"] == "EXCLUDED"
    assert ev_after["review_status"] == "UNRESOLVED"
    assert "품목 정의 충돌 객체는 CAD 자동 집계에 포함할 수 없습니다" in ev_after["review_reason"]

    # Verify reversed order in legacy storage produces identical safe result (no flip-flopping)
    tampered_reversed = list(reversed(tampered_mappings))
    (drawings_dir / "mappings.json").write_text(json.dumps(tampered_reversed), encoding="utf-8")
    t_rev = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    light_rev_row = next(r for r in t_rev["summary"] if r["item_code"] == "LIGHT")
    assert light_rev_row["quantity"] == 0
    assert light_rev_row["review_status"] == "UNRESOLVED"

    # Verify export headers
    csv_res = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/export/csv")
    assert csv_res.headers["X-Takeoff-Completeness"] == "provisional"
    assert int(csv_res.headers["X-Takeoff-Unresolved-Count"]) > 0

    xlsx_res = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/export/xlsx")
    assert xlsx_res.headers["X-Takeoff-Completeness"] == "provisional"
    assert int(xlsx_res.headers["X-Takeoff-Unresolved-Count"]) > 0


def test_manual_item_canonical_unit_integrity_adr_005(client: TestClient) -> None:
    """Issue #9 A01: Manual object canonical unit integrity, manual-only SET output, and conflict defenses."""
    proj = client.post("/api/projects", json={"name": "수동 기구 단위 정합성 검증"}).json()
    project_id = proj["id"]

    dxf_file = FIXTURES_DIR / "basic_devices.dxf"
    upload_res = client.post(
        f"/api/projects/{project_id}/drawings",
        files={"file": (dxf_file.name, dxf_file.read_bytes(), "application/octet-stream")},
    )
    drawing_id = upload_res.json()["drawing_id"]

    # 1. (MUNIT-E) Empty unit rejection in mappings
    empty_map_res = client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/mappings",
        json=[{"block_name": "LT-DL", "item_code": "LIGHT", "unit": ""}],
    )
    assert empty_map_res.status_code == 400
    assert "단위(unit)가 누락되었거나 공백입니다" in empty_map_res.json()["detail"]

    # 2. (MUNIT-E) Empty unit rejection in manual objects
    empty_mo_res = client.post(
        f"/api/projects/{project_id}/drawings/{drawing_id}/manual-objects",
        json={
            "id": "mo_empty",
            "drawing_id": drawing_id,
            "item_code": "CUSTOM",
            "unit": "   ",
            "x": 0.0,
            "y": 0.0,
        },
    )
    assert empty_mo_res.status_code == 400
    assert "단위(unit)가 누락되었거나 공백입니다" in empty_mo_res.json()["detail"]

    # 3. (MUNIT-A / BLOCKER-1) Unmapped manual-only non-EA rejected with 400 and 0 saved
    from eai_api.app import storage
    mo_set_unmapped = client.post(
        f"/api/projects/{project_id}/drawings/{drawing_id}/manual-objects",
        json={
            "id": "mo_panel_unmapped",
            "drawing_id": drawing_id,
            "item_code": "CUSTOM_SET",
            "item_name": "미매핑 분전반 세트",
            "unit": "SET",
            "x": 500.0,
            "y": 500.0,
        },
    )
    assert mo_set_unmapped.status_code == 400
    assert "P0 정책(M-UNIT-01)상 'EA' 단위만 허용됩니다" in mo_set_unmapped.json()["detail"]
    assert len(storage.get_manual_objects(project_id, drawing_id)) == 0, "No manual objects should be saved on 400"

    # 4. (MUNIT-B & MUNIT-H) Normal manual-only CUSTOM_EA with unit='EA' allowed and coordinates preserved
    mo_ea_valid = client.post(
        f"/api/projects/{project_id}/drawings/{drawing_id}/manual-objects",
        json={
            "id": "mo_custom_ea",
            "drawing_id": drawing_id,
            "item_code": "CUSTOM_EA",
            "item_name": "수동 추가 전등",
            "unit": "EA",
            "x": 123.4,
            "y": 567.8,
            "note": "현장 추가 조명",
        },
    )
    assert mo_ea_valid.status_code == 200

    t_b = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    custom_ea_row = next(r for r in t_b["summary"] if r["item_code"] == "CUSTOM_EA")
    assert custom_ea_row["quantity"] == 1
    assert custom_ea_row["unit"] == "EA"
    assert custom_ea_row["quantity_method"] == "MANUAL"

    custom_ea_ev = next(e for e in t_b["evidence"] if e["source_object_id"] == "mo_custom_ea")
    assert custom_ea_ev["x"] == 123.4
    assert custom_ea_ev["y"] == 567.8
    assert custom_ea_ev["inclusion_status"] == "MANUAL_ADDED"

    # 5. (MUNIT-C) Mapped OUTLET_SET with unit='SET' in mappings accepts manual SET additions
    client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/mappings",
        json=[
            {"block_name": "OT-2P", "item_code": "OUTLET_SET", "unit": "SET", "specification": "매입 2구 콘센트 세트"},
        ],
    )
    # Mismatched unit for mapped item rejected with 400
    mo_bad_unit = client.post(
        f"/api/projects/{project_id}/drawings/{drawing_id}/manual-objects",
        json={
            "id": "mo_outlet_set_bad",
            "drawing_id": drawing_id,
            "item_code": "OUTLET_SET",
            "unit": "EA",  # Mapped item is SET!
            "x": 300.0,
            "y": 300.0,
        },
    )
    assert mo_bad_unit.status_code == 400
    assert "도면 매핑 정규 단위('SET')와 일치하지 않습니다" in mo_bad_unit.json()["detail"]

    # Valid SET addition for mapped item
    mo_valid_set = client.post(
        f"/api/projects/{project_id}/drawings/{drawing_id}/manual-objects",
        json={
            "id": "mo_outlet_set_1",
            "drawing_id": drawing_id,
            "item_code": "OUTLET_SET",
            "unit": "SET",
            "x": 300.0,
            "y": 300.0,
            "note": "누락 콘센트 세트 추가",
        },
    )
    assert mo_valid_set.status_code == 200

    t_c = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    outlet_set_row = next(r for r in t_c["summary"] if r["item_code"] == "OUTLET_SET")
    # basic_devices.dxf has 4 OT-2P blocks + 1 manual = 5 SET!
    assert outlet_set_row["quantity"] == 5
    assert outlet_set_row["unit"] == "SET"
    assert outlet_set_row["quantity_method"] == "MIXED"

    # 6. (MUNIT-D / BLOCKER-2) Review override unit mismatch defense: EA manual reclassified to mapped SET code
    # API layer defense: PUT /reviews with mismatched unit returns 400 Bad Request
    rev_bad_res = client.put(
        f"/api/projects/{project_id}/drawings/{drawing_id}/reviews",
        json=[
            {
                "source_object_id": "mo_custom_ea",
                "inclusion_status": "INCLUDED",
                "review_status": "REVIEWED",
                "mapped_item_code": "OUTLET_SET",
                "reason": "단위 상충 재분류 시도",
            }
        ],
    )
    assert rev_bad_res.status_code == 400
    assert "단위('EA')와 재분류 대상 품목('OUTLET_SET')의 정규 단위('SET')가 상충합니다" in rev_bad_res.json()["detail"]

    # Storage fail-closed defense: Tampered/legacy reviews in storage must be isolated to EXCLUDED + UNRESOLVED
    drawings_dir = storage.projects_dir / project_id / "drawings" / drawing_id
    tampered_reviews = {
        "mo_custom_ea": {
            "source_object_id": "mo_custom_ea",
            "inclusion_status": "INCLUDED",
            "mapped_item_code": "OUTLET_SET",
            "review_status": "REVIEWED",
            "reason": "레거시 저장 상충 검수",
        }
    }
    (drawings_dir / "reviews.json").write_text(json.dumps(tampered_reviews), encoding="utf-8")

    t_d = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    # OUTLET_SET quantity must NOT include the conflicted EA object (still exactly 5 SET from 4 CAD + 1 valid manual SET)
    outlet_set_d = next(r for r in t_d["summary"] if r["item_code"] == "OUTLET_SET")
    assert outlet_set_d["quantity"] == 5
    assert outlet_set_d["unit"] == "SET"
    assert outlet_set_d["review_status"] == "UNRESOLVED", "Unresolved conflicting review must flag overall item status"

    conflicted_ev = next(e for e in t_d["evidence"] if e["source_object_id"] == "mo_custom_ea")
    assert conflicted_ev["inclusion_status"] == "EXCLUDED"
    assert conflicted_ev["review_status"] == "UNRESOLVED"
    assert "상충 (MUNIT-D)" in conflicted_ev["review_reason"]

    # Verify CSV and XLSX deep exports reflect consistent quantities and unresolved states
    csv_sum_res = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/export/csv?type=summary")
    csv_sum_text = csv_sum_res.content.decode("utf-8-sig")
    assert "OUTLET_SET" in csv_sum_text
    assert ",SET,5,MIXED,UNRESOLVED," in csv_sum_text

    csv_ev_res = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/export/csv?type=evidence")
    csv_ev_text = csv_ev_res.content.decode("utf-8-sig")
    assert "mo_custom_ea" in csv_ev_text
    assert "EXCLUDED,UNRESOLVED" in csv_ev_text

    xlsx_res = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/export/xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_res.content))
    ws_sum = wb["Takeoff Summary"]
    found_outlet_row = False
    for r in range(2, ws_sum.max_row + 1):
        if ws_sum.cell(row=r, column=1).value == "OUTLET_SET":
            assert ws_sum.cell(row=r, column=4).value == "SET"
            assert ws_sum.cell(row=r, column=5).value == 5
            assert ws_sum.cell(row=r, column=7).value == "UNRESOLVED"
            found_outlet_row = True
    assert found_outlet_row is True

    # 7. (MUNIT-F) Defense against legacy tampered storage mixing EA/SET or unmapped non-EA for manual items
    (drawings_dir / "reviews.json").write_text("{}", encoding="utf-8")
    tampered_manual = [
        {"id": "man_1", "drawing_id": drawing_id, "item_code": "CONF_MAN", "unit": "SET", "x": 10.0, "y": 10.0},
        {"id": "man_2", "drawing_id": drawing_id, "item_code": "CONF_MAN", "unit": "EA", "x": 20.0, "y": 20.0},
    ]
    (drawings_dir / "manual_objects.json").write_text(json.dumps(tampered_manual), encoding="utf-8")

    t_tampered = client.get(f"/api/projects/{project_id}/drawings/{drawing_id}/takeoff").json()
    conf_man_row = next(r for r in t_tampered["summary"] if r["item_code"] == "CONF_MAN")
    assert conf_man_row["quantity"] == 0, "Conflicting units on manual item MUST NOT be summed"
    assert conf_man_row["review_status"] == "UNRESOLVED"
    assert "품목 정의 충돌" in conf_man_row["notes"]
    assert conf_man_row["unit"] == "[상충 단위 충돌]", "Must not quietly display 0 EA or confirmed unit"




