"""Tests for PR A: PNG/JPG 2D raster image drawing upload, manual takeoff, and export (Issue #12)."""

from pathlib import Path

import pytest
from eai_api.app import app
from eai_api.image_parser import inspect_image
from eai_api.storage import Storage
from fastapi.testclient import TestClient

FIXTURE_IMAGE_PATH = Path("fixtures/images/floorplan_sample.png")


@pytest.fixture
def client(tmp_path: Path) -> TestClient:
    from eai_api import app as app_module

    app_module.storage = Storage(root=tmp_path)
    return TestClient(app)


def test_inspect_valid_png() -> None:
    raw = FIXTURE_IMAGE_PATH.read_bytes()
    meta = inspect_image(raw, "floorplan_sample.png")
    assert meta.file_type == "PNG"
    assert meta.width == 1000
    assert meta.height == 800
    assert meta.size_bytes == len(raw)
    assert len(meta.sha256) == 64


def test_inspect_invalid_or_corrupt_images() -> None:
    # 0 bytes
    with pytest.raises(ValueError, match="비어 있습니다"):
        inspect_image(b"", "empty.png")

    # Corrupt PNG header
    with pytest.raises(ValueError, match="지원하지 않거나 손상된"):
        inspect_image(b"\x89PNG\r\n\x1a\x00truncated", "corrupt.png")

    # Truncated PNG (too short for IHDR)
    with pytest.raises(ValueError, match="손상되었거나 읽을 수 없는"):
        inspect_image(b"\x89PNG\r\n\x1a\n1234", "short.png")

    # Not an image
    with pytest.raises(ValueError, match="지원하지 않거나 손상된"):
        inspect_image(b"SECTION\n2\nHEADER\n0\nENDSEC\n0\nEOF\n", "fake.png")


def test_upload_image_drawing(client: TestClient) -> None:
    # 1. Create project
    proj_res = client.post("/api/projects", json={"name": "이미지 도면 테스트 프로젝트"})
    assert proj_res.status_code == 200
    proj_id = proj_res.json()["id"]

    # 2. Upload PNG
    raw_bytes = FIXTURE_IMAGE_PATH.read_bytes()
    upload_res = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("floorplan_sample.png", raw_bytes, "image/png")},
    )
    assert upload_res.status_code == 200
    upload_data = upload_res.json()
    drawing_id = upload_data["drawing_id"]
    assert upload_data["source_type"] == "IMAGE"
    assert upload_data["file_type"] == "PNG"
    assert upload_data["coordinate_space"] == "IMAGE_PIXEL"
    assert upload_data["dimensions"] == {"width": 1000, "height": 800}
    assert upload_data["objects_count"] == 0
    assert upload_data["warnings_count"] >= 1

    # 3. Get drawing detail
    dwg_res = client.get(f"/api/projects/{proj_id}/drawings/{drawing_id}")
    assert dwg_res.status_code == 200
    dwg = dwg_res.json()
    assert dwg["source_type"] == "IMAGE"
    assert dwg["file_type"] == "PNG"
    assert dwg["coordinate_space"] == "IMAGE_PIXEL"
    assert dwg["extents"]["min"] == [0.0, 0.0]
    assert dwg["extents"]["max"] == [1000.0, 800.0]
    assert dwg["dxf"]["units_name"] == "PIXEL"
    # Informational warning that manual EA marker mode is active and scale is uncalibrated
    assert "픽셀 좌표계(IMAGE_PIXEL)" in dwg["warnings"][0]["message"]
    assert "FR-016" in dwg["warnings"][0]["message"]

    # 4. Image serving endpoint
    img_res = client.get(f"/api/projects/{proj_id}/drawings/{drawing_id}/image")
    assert img_res.status_code == 200
    assert img_res.headers["content-type"] == "image/png"
    assert img_res.content == raw_bytes

    # 5. Raw download endpoint
    raw_res = client.get(f"/api/projects/{proj_id}/drawings/{drawing_id}/raw")
    assert raw_res.status_code == 200
    assert raw_res.headers["content-type"] == "image/png"
    assert "floorplan_sample.png" in raw_res.headers["content-disposition"]
    assert raw_res.content == raw_bytes


def test_reject_dwg_upload(client: TestClient) -> None:
    proj_res = client.post("/api/projects", json={"name": "DWG 거부 테스트"})
    proj_id = proj_res.json()["id"]

    res = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("drawing.dwg", b"AC1032binarydwgdata", "application/octet-stream")},
    )
    assert res.status_code == 400
    assert "DWG 직접 업로드는 지원하지 않습니다" in res.json()["detail"]


def test_image_manual_takeoff_workflow_and_exports(client: TestClient) -> None:
    proj_res = client.post("/api/projects", json={"name": "수동 적산 프로젝트"})
    proj_id = proj_res.json()["id"]

    raw_bytes = FIXTURE_IMAGE_PATH.read_bytes()
    upload_res = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("floorplan_sample.png", raw_bytes, "image/png")},
    )
    drawing_id = upload_res.json()["drawing_id"]

    # Initial takeoff: 0 items
    t0_res = client.get(f"/api/projects/{proj_id}/drawings/{drawing_id}/takeoff")
    assert t0_res.status_code == 200
    t0 = t0_res.json()
    assert len(t0["summary"]) == 0
    assert len(t0["evidence"]) == 0

    # Add 4 LIGHT markers, 2 OUTLET markers, 2 SWITCH markers at fixture pixel coords
    markers = [
        {"id": "man_lt_1", "drawing_id": drawing_id, "x": 250, "y": 250, "item_code": "LIGHT", "item_name": "다운라이트", "unit": "EA"},
        {"id": "man_lt_2", "drawing_id": drawing_id, "x": 750, "y": 250, "item_code": "LIGHT", "item_name": "다운라이트", "unit": "EA"},
        {"id": "man_lt_3", "drawing_id": drawing_id, "x": 250, "y": 550, "item_code": "LIGHT", "item_name": "다운라이트", "unit": "EA"},
        {"id": "man_lt_4", "drawing_id": drawing_id, "x": 750, "y": 550, "item_code": "LIGHT", "item_name": "다운라이트", "unit": "EA"},
        {"id": "man_ot_1", "drawing_id": drawing_id, "x": 120, "y": 250, "item_code": "OUTLET", "item_name": "2구 콘센트", "unit": "EA"},
        {"id": "man_ot_2", "drawing_id": drawing_id, "x": 880, "y": 250, "item_code": "OUTLET", "item_name": "2구 콘센트", "unit": "EA"},
        {"id": "man_sw_1", "drawing_id": drawing_id, "x": 460, "y": 380, "item_code": "SWITCH", "item_name": "1구 스위치", "unit": "EA"},
        {"id": "man_sw_2", "drawing_id": drawing_id, "x": 540, "y": 380, "item_code": "SWITCH", "item_name": "1구 스위치", "unit": "EA"},
    ]

    for m in markers:
        add_res = client.post(
            f"/api/projects/{proj_id}/drawings/{drawing_id}/manual-objects",
            json=m,
        )
        assert add_res.status_code == 200

    # Verify P0 M-UNIT-01 defense: non-EA manual item must be rejected
    bad_res = client.post(
        f"/api/projects/{proj_id}/drawings/{drawing_id}/manual-objects",
        json={"id": "bad_unit", "drawing_id": drawing_id, "x": 100, "y": 100, "item_code": "CABLE", "unit": "M"},
    )
    assert bad_res.status_code == 400
    assert "M-UNIT-01" in bad_res.json()["detail"]

    # Re-fetch takeoff
    t1_res = client.get(f"/api/projects/{proj_id}/drawings/{drawing_id}/takeoff")
    assert t1_res.status_code == 200
    t1 = t1_res.json()

    # Summary checks: 3 items (LIGHT: 4 EA, OUTLET: 2 EA, SWITCH: 2 EA)
    sum_by_code = {row["item_code"]: row for row in t1["summary"]}
    assert set(sum_by_code.keys()) == {"LIGHT", "OUTLET", "SWITCH"}

    assert sum_by_code["LIGHT"]["quantity"] == 4
    assert sum_by_code["LIGHT"]["unit"] == "EA"
    assert sum_by_code["LIGHT"]["quantity_method"] == "MANUAL"

    assert sum_by_code["OUTLET"]["quantity"] == 2
    assert sum_by_code["OUTLET"]["unit"] == "EA"
    assert sum_by_code["OUTLET"]["quantity_method"] == "MANUAL"

    assert sum_by_code["SWITCH"]["quantity"] == 2
    assert sum_by_code["SWITCH"]["unit"] == "EA"
    assert sum_by_code["SWITCH"]["quantity_method"] == "MANUAL"

    # Evidence checks: 8 rows, entity_type="MANUAL_POINT", pixel coords preserved
    assert len(t1["evidence"]) == 8
    ev_by_id = {row["source_object_id"]: row for row in t1["evidence"]}
    assert ev_by_id["man_lt_1"]["x"] == 250.0
    assert ev_by_id["man_lt_1"]["y"] == 250.0
    assert ev_by_id["man_lt_1"]["entity_type"] == "MANUAL_POINT"
    assert ev_by_id["man_lt_1"]["block_name"] == "[MANUAL]"
    assert ev_by_id["man_lt_1"]["inclusion_status"] == "MANUAL_ADDED"

    # Export CSV Summary
    csv_sum = client.get(f"/api/projects/{proj_id}/drawings/{drawing_id}/export/csv?type=summary")
    assert csv_sum.status_code == 200
    text_sum = csv_sum.content.decode("utf-8-sig")
    assert "LIGHT,다운라이트,,EA,4,MANUAL" in text_sum
    assert "OUTLET,2구 콘센트,,EA,2,MANUAL" in text_sum
    assert "SWITCH,1구 스위치,,EA,2,MANUAL" in text_sum

    # Export CSV Evidence
    csv_ev = client.get(f"/api/projects/{proj_id}/drawings/{drawing_id}/export/csv?type=evidence")
    assert csv_ev.status_code == 200
    text_ev = csv_ev.content.decode("utf-8-sig")
    assert "man_lt_1" in text_ev
    assert "MANUAL_POINT" in text_ev
    assert "250" in text_ev

    # Export XLSX
    xlsx_res = client.get(f"/api/projects/{proj_id}/drawings/{drawing_id}/export/xlsx")
    assert xlsx_res.status_code == 200
    assert len(xlsx_res.content) > 1000

    # Delete one marker (man_lt_4)
    del_res = client.delete(f"/api/projects/{proj_id}/drawings/{drawing_id}/manual-objects/man_lt_4")
    assert del_res.status_code == 200

    # Recheck takeoff: LIGHT should now be 3 EA
    t2_res = client.get(f"/api/projects/{proj_id}/drawings/{drawing_id}/takeoff")
    t2 = t2_res.json()
    sum_by_code_2 = {row["item_code"]: row for row in t2["summary"]}
    assert sum_by_code_2["LIGHT"]["quantity"] == 3
    assert len(t2["evidence"]) == 7


def test_upload_and_stream_valid_jpeg(client: TestClient) -> None:
    """Test valid JPEG upload, metadata, image streaming, and raw download."""
    jpeg_path = Path("fixtures/images/floorplan_sample.jpg")
    raw_bytes = jpeg_path.read_bytes()

    # Create project
    proj_res = client.post("/api/projects", json={"name": "JPEG 업로드 테스트"})
    assert proj_res.status_code == 200
    proj_id = proj_res.json()["id"]

    # Upload JPEG
    upload_res = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("floorplan_sample.jpg", raw_bytes, "image/jpeg")},
    )
    assert upload_res.status_code == 200
    data = upload_res.json()
    assert data["source_type"] == "IMAGE"
    assert data["file_type"] == "JPG"
    assert data["coordinate_space"] == "IMAGE_PIXEL"
    assert data["dimensions"] == {"width": 1000, "height": 800}

    drawing_id = data["drawing_id"]

    # Image stream check
    stream_res = client.get(f"/api/projects/{proj_id}/drawings/{drawing_id}/image")
    assert stream_res.status_code == 200
    assert stream_res.headers["content-type"] == "image/jpeg"
    assert stream_res.content == raw_bytes

    # Raw download check
    raw_res = client.get(f"/api/projects/{proj_id}/drawings/{drawing_id}/raw")
    assert raw_res.status_code == 200
    assert raw_res.headers["content-type"] == "image/jpeg"
    assert raw_res.content == raw_bytes


def test_corrupted_image_upload_rejected(client: TestClient) -> None:
    """Test corrupted or truncated PNG and JPEG are rejected with 400 Bad Request."""
    proj_res = client.post("/api/projects", json={"name": "손상 이미지 거부 테스트"})
    proj_id = proj_res.json()["id"]

    # 1. Fake PNG payload with valid magic bytes but corrupted data
    fake_png = b"\x89PNG\r\n\x1a\n" + b"\x00" * 200
    res1 = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("fake.png", fake_png, "image/png")},
    )
    assert res1.status_code == 400
    assert "손상" in res1.json()["detail"] or "지원하지 않거나" in res1.json()["detail"]

    # 2. Truncated PNG
    valid_png = FIXTURE_IMAGE_PATH.read_bytes()
    truncated_png = valid_png[:300]
    res2 = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("truncated.png", truncated_png, "image/png")},
    )
    assert res2.status_code == 400
    assert "손상" in res2.json()["detail"] or "truncated" in res2.json()["detail"]

    # 3. Truncated JPEG
    fake_jpg = b"\xff\xd8\xff\xe0" + b"\x00" * 50
    res3 = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("fake.jpg", fake_jpg, "image/jpeg")},
    )
    assert res3.status_code == 400
    assert "손상" in res3.json()["detail"] or "지원하지 않거나" in res3.json()["detail"]


def test_exif_orientation_rejected(client: TestClient) -> None:
    """Test JPEG with EXIF orientation (e.g. 6) is rejected with clear 400 diagnostic."""
    import io

    from PIL import Image

    # Create in-memory JPEG with EXIF orientation 6 (rotated 90 deg)
    im = Image.new("RGB", (200, 100), color=(200, 200, 200))
    exif = im.getexif()
    exif[0x0112] = 6  # Orientation: 6
    buf = io.BytesIO()
    im.save(buf, format="JPEG", exif=exif)
    jpeg_with_exif = buf.getvalue()

    proj_res = client.post("/api/projects", json={"name": "EXIF 회전 테스트"})
    proj_id = proj_res.json()["id"]

    res = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("rotated_photo.jpg", jpeg_with_exif, "image/jpeg")},
    )
    assert res.status_code == 400
    detail = res.json()["detail"]
    assert "EXIF 회전 메타데이터" in detail
    assert "Orientation: 6" in detail
    assert "IMAGE_PIXEL" in detail


def test_manual_marker_bounds_on_image_drawing(client: TestClient) -> None:
    """Test manual marker boundary check on IMAGE drawings [0, w) x [0, h) half-open interval."""
    proj_res = client.post("/api/projects", json={"name": "좌표 경계 테스트"})
    proj_id = proj_res.json()["id"]

    raw_bytes = FIXTURE_IMAGE_PATH.read_bytes()
    upload_res = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("floorplan_sample.png", raw_bytes, "image/png")},
    )
    drawing_id = upload_res.json()["drawing_id"]

    # 1. Negative coordinates should be rejected with 400
    neg_x = client.post(
        f"/api/projects/{proj_id}/drawings/{drawing_id}/manual-objects",
        json={"id": "m_neg_x", "drawing_id": drawing_id, "x": -0.1, "y": 100.0, "item_code": "LIGHT", "unit": "EA"},
    )
    assert neg_x.status_code == 400
    assert "벗어났습니다" in neg_x.json()["detail"]

    neg_y = client.post(
        f"/api/projects/{proj_id}/drawings/{drawing_id}/manual-objects",
        json={"id": "m_neg_y", "drawing_id": drawing_id, "x": 100.0, "y": -0.1, "item_code": "LIGHT", "unit": "EA"},
    )
    assert neg_y.status_code == 400
    assert "벗어났습니다" in neg_y.json()["detail"]

    # 2. Boundary edge at width/height (1000.0, 800.0) must be rejected with 400 (half-open interval)
    edge_x = client.post(
        f"/api/projects/{proj_id}/drawings/{drawing_id}/manual-objects",
        json={"id": "m_edge_x", "drawing_id": drawing_id, "x": 1000.0, "y": 0.0, "item_code": "LIGHT", "unit": "EA"},
    )
    assert edge_x.status_code == 400
    assert "벗어났습니다" in edge_x.json()["detail"]

    edge_y = client.post(
        f"/api/projects/{proj_id}/drawings/{drawing_id}/manual-objects",
        json={"id": "m_edge_y", "drawing_id": drawing_id, "x": 0.0, "y": 800.0, "item_code": "LIGHT", "unit": "EA"},
    )
    assert edge_y.status_code == 400
    assert "벗어났습니다" in edge_y.json()["detail"]

    edge_corner = client.post(
        f"/api/projects/{proj_id}/drawings/{drawing_id}/manual-objects",
        json={"id": "m_edge_corner", "drawing_id": drawing_id, "x": 1000.0, "y": 800.0, "item_code": "LIGHT", "unit": "EA"},
    )
    assert edge_corner.status_code == 400
    assert "벗어났습니다" in edge_corner.json()["detail"]

    # 3. Beyond Width and Height
    over_res = client.post(
        f"/api/projects/{proj_id}/drawings/{drawing_id}/manual-objects",
        json={"id": "m_over", "drawing_id": drawing_id, "x": 1001.0, "y": 801.0, "item_code": "LIGHT", "unit": "EA"},
    )
    assert over_res.status_code == 400
    assert "벗어났습니다" in over_res.json()["detail"]

    # 4. Non-finite values (NaN / Inf) should be rejected
    nan_res = client.post(
        f"/api/projects/{proj_id}/drawings/{drawing_id}/manual-objects",
        json={"id": "m_nan", "drawing_id": drawing_id, "x": "NaN", "y": 100.0, "item_code": "LIGHT", "unit": "EA"},
    )
    assert nan_res.status_code in (400, 422)

    # 5. Valid bounds within [0, 1000) x [0, 800):
    # (0, 0), (999, 0), (0, 799), (999, 799), (999.9, 799.9) must succeed (200)
    valid_points = [
        ("p_origin", 0.0, 0.0),
        ("p_top_right_max_int", 999.0, 0.0),
        ("p_bottom_left_max_int", 0.0, 799.0),
        ("p_bottom_right_max_int", 999.0, 799.0),
        ("p_bottom_right_float", 999.9, 799.9),
    ]
    for pid, px, py in valid_points:
        p_res = client.post(
            f"/api/projects/{proj_id}/drawings/{drawing_id}/manual-objects",
            json={"id": pid, "drawing_id": drawing_id, "x": px, "y": py, "item_code": "LIGHT", "unit": "EA"},
        )
        assert p_res.status_code == 200, f"Failed on valid point ({px}, {py}): {p_res.text}"

    # Verify takeoff has 5 items
    t_res = client.get(f"/api/projects/{proj_id}/drawings/{drawing_id}/takeoff")
    assert t_res.status_code == 200
    assert t_res.json()["summary"][0]["quantity"] == 5


def test_image_drawing_img03_unit_security_enforcement(client: TestClient) -> None:
    """Test IMG-03 / M-UNIT-01 security: Image drawings strictly enforce EA-only units across API layers."""
    proj_res = client.post("/api/projects", json={"name": "IMG-03 보안 테스트"})
    proj_id = proj_res.json()["id"]

    raw_bytes = FIXTURE_IMAGE_PATH.read_bytes()
    upload_res = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("floorplan_sample.png", raw_bytes, "image/png")},
    )
    drawing_id = upload_res.json()["drawing_id"]

    # 1. Attempt PUT /mappings with non-EA unit on IMAGE drawing must be rejected with 400
    bad_map_res = client.put(
        f"/api/projects/{proj_id}/drawings/{drawing_id}/mappings",
        json=[
            {
                "block_name": "TEST_BLOCK",
                "item_code": "PANEL_SET",
                "item_name": "분전반",
                "unit": "SET",
            }
        ],
    )
    assert bad_map_res.status_code == 400
    assert "IMG-03" in bad_map_res.json()["detail"] or "EA" in bad_map_res.json()["detail"]

    # Valid mapping with EA unit succeeds
    ok_map_res = client.put(
        f"/api/projects/{proj_id}/drawings/{drawing_id}/mappings",
        json=[
            {
                "block_name": "TEST_BLOCK",
                "item_code": "LIGHT_EA",
                "item_name": "LED조명",
                "unit": "EA",
            }
        ],
    )
    assert ok_map_res.status_code == 200

    # 2. Attempt POST /manual-objects with non-EA unit must be rejected with 400
    bad_man_res = client.post(
        f"/api/projects/{proj_id}/drawings/{drawing_id}/manual-objects",
        json={"id": "m_bad_unit", "drawing_id": drawing_id, "x": 100.0, "y": 100.0, "item_code": "LIGHT_EA", "unit": "SET"},
    )
    assert bad_man_res.status_code == 400
    assert "IMG-03" in bad_man_res.json()["detail"]

    # Add valid manual object with EA unit
    ok_man_res = client.post(
        f"/api/projects/{proj_id}/drawings/{drawing_id}/manual-objects",
        json={"id": "m_ok_1", "drawing_id": drawing_id, "x": 100.0, "y": 100.0, "item_code": "LIGHT_EA", "unit": "EA"},
    )
    assert ok_man_res.status_code == 200

    # 3. Attempt PUT /reviews to reclassify with non-EA target unit:
    # First, verify reclassifying to EXCLUDE or UNMAPPED is allowed
    rev_ok = client.put(
        f"/api/projects/{proj_id}/drawings/{drawing_id}/reviews",
        json=[
            {
                "source_object_id": "m_ok_1",
                "mapped_item_code": "EXCLUDE",
                "inclusion_status": "EXCLUDED",
                "review_status": "REVIEWED",
                "reason": "검수 제외",
            }
        ],
    )
    assert rev_ok.status_code == 200

    # Reclassify back to LIGHT_EA
    rev_ok2 = client.put(
        f"/api/projects/{proj_id}/drawings/{drawing_id}/reviews",
        json=[
            {
                "source_object_id": "m_ok_1",
                "mapped_item_code": "LIGHT_EA",
                "inclusion_status": "INCLUDED",
                "review_status": "REVIEWED",
                "reason": "정상 포함",
            }
        ],
    )
    assert rev_ok2.status_code == 200

    # 4. Fail-closed defense: Tampered object with non-EA unit directly in storage
    from eai_api import app as app_module
    from eai_api.models import ManualObjectItem

    app_module.storage.save_manual_objects(
        proj_id,
        drawing_id,
        [
            ManualObjectItem(
                id="m_ok_1",
                drawing_id=drawing_id,
                x=100.0,
                y=100.0,
                item_code="LIGHT_EA",
                item_name="LED조명",
                unit="EA",
            ),
            ManualObjectItem(
                id="m_tampered_non_ea",
                drawing_id=drawing_id,
                x=200.0,
                y=200.0,
                item_code="TAMPERED_CABLE",
                item_name="케이블",
                unit="M",
            ),
        ],
    )

    t_res = client.get(f"/api/projects/{proj_id}/drawings/{drawing_id}/takeoff")
    assert t_res.status_code == 200
    t_data = t_res.json()

    # The tampered non-EA object must be EXCLUDED and UNRESOLVED in evidence
    tampered_ev = next(e for e in t_data["evidence"] if e["source_object_id"] == "m_tampered_non_ea")
    assert tampered_ev["inclusion_status"] == "EXCLUDED"
    assert tampered_ev["review_status"] == "UNRESOLVED"
    assert "비EA 불허" in tampered_ev["review_reason"]

    # In summary rows, TAMPERED_CABLE must have quantity == 0 and UNRESOLVED
    sum_by_code = {s["item_code"]: s for s in t_data["summary"]}
    assert sum_by_code["TAMPERED_CABLE"]["quantity"] == 0
    assert sum_by_code["TAMPERED_CABLE"]["review_status"] == "UNRESOLVED"

    # LIGHT_EA must have quantity == 1 and unit == "EA"
    assert sum_by_code["LIGHT_EA"]["quantity"] == 1
    assert sum_by_code["LIGHT_EA"]["unit"] == "EA"

    # 5. Export CSV & XLSX safety: CSV Summary should have quantity 0 for TAMPERED_CABLE
    csv_sum = client.get(f"/api/projects/{proj_id}/drawings/{drawing_id}/export/csv?type=summary")
    assert csv_sum.status_code == 200
    csv_text = csv_sum.content.decode("utf-8-sig")
    assert "TAMPERED_CABLE" in csv_text
    assert ",0,MANUAL,UNRESOLVED," in csv_text
    assert "[상충 단위 충돌]" in csv_text
    assert "LIGHT_EA,LED조명,,EA,1,MANUAL" in csv_text

    # Also verify XLSX export
    xlsx_res = client.get(f"/api/projects/{proj_id}/drawings/{drawing_id}/export/xlsx")
    assert xlsx_res.status_code == 200
    assert len(xlsx_res.content) > 1000


def test_manual_post_drawing_id_mismatch_rejected(client: TestClient) -> None:
    """Test traceability guard: POST /manual-objects with body drawing_id mismatching URL path is rejected with 400."""
    proj_res = client.post("/api/projects", json={"name": "drawing_id 일치 검증"})
    proj_id = proj_res.json()["id"]

    raw_bytes = FIXTURE_IMAGE_PATH.read_bytes()
    upload_res = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("floorplan_sample.png", raw_bytes, "image/png")},
    )
    drawing_id = upload_res.json()["drawing_id"]

    mismatch_res = client.post(
        f"/api/projects/{proj_id}/drawings/{drawing_id}/manual-objects",
        json={
            "id": "m_mismatch",
            "drawing_id": "different_dwg_id",
            "x": 100.0,
            "y": 100.0,
            "item_code": "LIGHT",
            "unit": "EA",
        },
    )
    assert mismatch_res.status_code == 400
    assert "drawing_id" in mismatch_res.json()["detail"]


def test_legacy_or_tampered_off_image_manual_markers_fail_closed(client: TestClient) -> None:
    """Test A04/A01 HIGH defense: legacy/tampered off-image manual markers in storage fail-closed at GET takeoff."""
    import csv
    import io

    import openpyxl
    from eai_api import app as app_module
    from eai_api.models import EvidenceReviewItem, InclusionStatus, ManualObjectItem, ReviewStatus

    proj_res = client.post("/api/projects", json={"name": "레거시 오프이미지 마커 방어 테스트"})
    proj_id = proj_res.json()["id"]

    raw_bytes = FIXTURE_IMAGE_PATH.read_bytes()
    upload_res = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("floorplan_sample.png", raw_bytes, "image/png")},
    )
    drawing_id = upload_res.json()["drawing_id"]

    # Direct storage injection of legacy or manually tampered markers:
    # 1 valid marker at boundary (999.0, 799.0)
    # 5 off-image / non-finite markers: (1000.0, 0.0), (0.0, 800.0), (-1.0, 1.0), (NaN, 100.0), (100.0, Inf)
    app_module.storage.save_manual_objects(
        proj_id,
        drawing_id,
        [
            ManualObjectItem.model_construct(id="m_valid", drawing_id=drawing_id, x=999.0, y=799.0, item_code="LIGHT", item_name="전등", unit="EA"),
            ManualObjectItem.model_construct(id="m_off_x", drawing_id=drawing_id, x=1000.0, y=0.0, item_code="LIGHT", item_name="전등", unit="EA"),
            ManualObjectItem.model_construct(id="m_off_y", drawing_id=drawing_id, x=0.0, y=800.0, item_code="LIGHT", item_name="전등", unit="EA"),
            ManualObjectItem.model_construct(id="m_neg", drawing_id=drawing_id, x=-1.0, y=1.0, item_code="LIGHT", item_name="전등", unit="EA"),
            ManualObjectItem.model_construct(id="m_nan", drawing_id=drawing_id, x=float("nan"), y=100.0, item_code="LIGHT", item_name="전등", unit="EA"),
            ManualObjectItem.model_construct(id="m_inf", drawing_id=drawing_id, x=100.0, y=float("inf"), item_code="LIGHT", item_name="전등", unit="EA"),
        ],
    )

    # Tampered review override trying to force inclusion on m_off_x
    app_module.storage.save_reviews(
        proj_id,
        drawing_id,
        {
            "m_off_x": EvidenceReviewItem(
                source_object_id="m_off_x",
                inclusion_status=InclusionStatus.INCLUDED,
                review_status=ReviewStatus.REVIEWED,
                reason="변조된 강제 포함 시도",
            )
        },
    )

    # 1. Verify GET /takeoff response
    t_res = client.get(f"/api/projects/{proj_id}/drawings/{drawing_id}/takeoff")
    assert t_res.status_code == 200
    t_data = t_res.json()

    # Summary: LIGHT quantity must be strictly 1 (preserving only m_valid, excluding all 5 off-image markers)
    sum_by_code = {s["item_code"]: s for s in t_data["summary"]}
    assert "LIGHT" in sum_by_code
    assert sum_by_code["LIGHT"]["quantity"] == 1
    assert sum_by_code["LIGHT"]["unit"] == "EA"
    assert sum_by_code["LIGHT"]["review_status"] == "UNRESOLVED"
    assert "미해결 객체" in sum_by_code["LIGHT"]["notes"]

    # Provisional and unresolved counts
    assert t_data["is_provisional"] is True
    assert t_data["unresolved_count"] >= 5

    # Evidence: check each marker
    ev_by_id = {e["source_object_id"]: e for e in t_data["evidence"]}
    assert ev_by_id["m_valid"]["inclusion_status"] == "MANUAL_ADDED"
    assert ev_by_id["m_valid"]["review_status"] == "REVIEWED"
    assert ev_by_id["m_valid"]["x"] == 999.0
    assert ev_by_id["m_valid"]["y"] == 799.0

    for off_id in ("m_off_x", "m_off_y", "m_neg", "m_nan", "m_inf"):
        assert off_id in ev_by_id
        ev = ev_by_id[off_id]
        assert ev["inclusion_status"] == "EXCLUDED", f"{off_id} must be EXCLUDED"
        assert ev["review_status"] == "UNRESOLVED", f"{off_id} must be UNRESOLVED"
        assert "이탈/비유한 좌표" in ev["review_reason"]
        assert "IMG-02" in ev["review_reason"]

    # 2. Verify CSV Summary with csv.reader
    csv_sum_res = client.get(f"/api/projects/{proj_id}/drawings/{drawing_id}/export/csv?type=summary")
    assert csv_sum_res.status_code == 200
    sum_reader = list(csv.reader(io.StringIO(csv_sum_res.content.decode("utf-8-sig"))))
    sum_header = sum_reader[0]
    qty_col = sum_header.index("quantity")
    status_col = sum_header.index("review_status")
    code_col = sum_header.index("item_code")

    light_row = next(r for r in sum_reader[1:] if r[code_col] == "LIGHT")
    assert light_row[qty_col] == "1"
    assert light_row[status_col] == "UNRESOLVED"

    # 3. Verify CSV Evidence with csv.reader
    csv_ev_res = client.get(f"/api/projects/{proj_id}/drawings/{drawing_id}/export/csv?type=evidence")
    assert csv_ev_res.status_code == 200
    ev_reader = list(csv.reader(io.StringIO(csv_ev_res.content.decode("utf-8-sig"))))
    ev_header = ev_reader[0]
    id_col = ev_header.index("source_object_id")
    inc_col = ev_header.index("inclusion_status")
    rev_col = ev_header.index("review_status")

    ev_rows_by_id = {r[id_col]: r for r in ev_reader[1:]}
    assert ev_rows_by_id["m_valid"][inc_col] == "MANUAL_ADDED"
    assert ev_rows_by_id["m_valid"][rev_col] == "REVIEWED"
    for off_id in ("m_off_x", "m_off_y", "m_neg", "m_nan", "m_inf"):
        assert ev_rows_by_id[off_id][inc_col] == "EXCLUDED"
        assert ev_rows_by_id[off_id][rev_col] == "UNRESOLVED"

    # 4. Verify XLSX with openpyxl cell-by-cell
    xlsx_res = client.get(f"/api/projects/{proj_id}/drawings/{drawing_id}/export/xlsx")
    assert xlsx_res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_res.content))
    assert "Takeoff Summary" in wb.sheetnames
    assert "Takeoff Evidence" in wb.sheetnames

    # Check Summary sheet: row for LIGHT has quantity == 1
    ws_sum = wb["Takeoff Summary"]
    found_light_sum = False
    for row in ws_sum.iter_rows(min_row=2, values_only=True):
        if row[0] == "LIGHT":
            found_light_sum = True
            assert row[4] == 1  # quantity
            assert row[6] == "UNRESOLVED"  # review_status
    assert found_light_sum is True

    # Check Evidence sheet: verify cells for invalid markers have EXCLUDED and UNRESOLVED
    ws_ev = wb["Takeoff Evidence"]
    ev_xlsx_by_id = {}
    for row in ws_ev.iter_rows(min_row=2, values_only=True):
        ev_xlsx_by_id[row[0]] = {"inclusion": row[9], "review": row[10]}

    assert ev_xlsx_by_id["m_valid"]["inclusion"] == "MANUAL_ADDED"
    assert ev_xlsx_by_id["m_valid"]["review"] == "REVIEWED"
    for off_id in ("m_off_x", "m_off_y", "m_neg", "m_nan", "m_inf"):
        assert ev_xlsx_by_id[off_id]["inclusion"] == "EXCLUDED"
        assert ev_xlsx_by_id[off_id]["review"] == "UNRESOLVED"

    # 5. Immutability verification: raw image file on disk remains 100% identical
    raw_res = client.get(f"/api/projects/{proj_id}/drawings/{drawing_id}/raw")
    assert raw_res.status_code == 200
    assert raw_res.content == raw_bytes


def test_dxf_wcs_negative_coordinates_preserved_and_counted(client: TestClient) -> None:
    """Test DXF CAD drawings preserve and correctly count negative WCS manual coordinates."""
    proj_res = client.post("/api/projects", json={"name": "DXF 음수 WCS 좌표 테스트"})
    proj_id = proj_res.json()["id"]

    dxf_path = Path("fixtures/dxf/basic_devices.dxf")
    upload_res = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("basic_devices.dxf", dxf_path.read_bytes(), "application/octet-stream")},
    )
    assert upload_res.status_code == 200
    drawing_id = upload_res.json()["drawing_id"]

    # Add manual marker with negative WCS coordinates on CAD drawing
    neg_cad = client.post(
        f"/api/projects/{proj_id}/drawings/{drawing_id}/manual-objects",
        json={"id": "m_cad_neg", "drawing_id": drawing_id, "x": -150.0, "y": -300.0, "item_code": "LIGHT_EXT", "unit": "EA"},
    )
    assert neg_cad.status_code == 200

    t_res = client.get(f"/api/projects/{proj_id}/drawings/{drawing_id}/takeoff")
    assert t_res.status_code == 200
    t_data = t_res.json()

    # Negative coordinate on DXF must be included and counted
    sum_by_code = {s["item_code"]: s for s in t_data["summary"]}
    assert "LIGHT_EXT" in sum_by_code
    assert sum_by_code["LIGHT_EXT"]["quantity"] == 1
    assert sum_by_code["LIGHT_EXT"]["review_status"] == "REVIEWED"

    ev = next(e for e in t_data["evidence"] if e["source_object_id"] == "m_cad_neg")
    assert ev["inclusion_status"] == "MANUAL_ADDED"
    assert ev["x"] == -150.0
    assert ev["y"] == -300.0

