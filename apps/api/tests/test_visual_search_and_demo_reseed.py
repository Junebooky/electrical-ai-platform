"""Tests for visual symbol template search engine and demo reseed idempotency (Issue #29)."""

from __future__ import annotations

import csv
import io
from pathlib import Path

import openpyxl
import pytest
from eai_api.app import app
from eai_api.models import (
    CandidateBulkReviewRequest,
    CandidateReviewUpdate,
    CandidateStatus,
    ItemMappingRule,
    SymbolCandidate,
)
from eai_api.storage import Storage
from fastapi.testclient import TestClient

DXF_FIXTURE_PATH = Path("fixtures/dxf/basic_devices.dxf")
PDF_FIXTURE_PATH = Path("fixtures/pdf/pdf_c_vector_features.pdf")
IMAGE_FIXTURE_PATH = Path("fixtures/images/floorplan_sample.png")


@pytest.fixture
def client(tmp_path: Path) -> TestClient:
    from eai_api import app as app_module

    app_module.storage = Storage(root=tmp_path)
    return TestClient(app)


def test_demo_reseed_idempotency_and_preservation(client: TestClient) -> None:
    """Verify P1-DEMO-RESEED: calling /api/demo/seed repeatedly does NOT accumulate drawings and preserves user reviews."""
    from eai_api import app as app_module

    # 1. Initial seed
    res1 = client.post("/api/demo/seed")
    assert res1.status_code == 200
    data1 = res1.json()
    proj_id = data1["project_id"]
    drawings1 = data1["drawings"]
    assert len(drawings1) == 3

    # Check total drawings in storage
    all_drawings_init = app_module.storage.list_drawings(proj_id)
    assert len(all_drawings_init) == 3

    # 2. Simulate user reviewing candidates on the image drawing
    img_dwg = next(d for d in drawings1 if d["source_type"] == "IMAGE")
    img_dwg_id = img_dwg["drawing_id"]
    cands_res = client.get(f"/api/projects/{proj_id}/drawings/{img_dwg_id}/candidates")
    cands = cands_res.json()["candidates"]
    assert len(cands) >= 2

    # Accept the first candidate and reject the second
    target_accept_id = cands[0]["candidate_id"]
    target_reject_id = cands[1]["candidate_id"]

    bulk_req = CandidateBulkReviewRequest(
        updates=[
            CandidateReviewUpdate(candidate_id=target_accept_id, status=CandidateStatus.ACCEPTED),
            CandidateReviewUpdate(candidate_id=target_reject_id, status=CandidateStatus.REJECTED),
        ]
    )
    rev_res = client.put(
        f"/api/projects/{proj_id}/drawings/{img_dwg_id}/candidates/review",
        json=bulk_req.model_dump(mode="json"),
    )
    assert rev_res.status_code == 200

    # 3. Second seed call with default reset=False (Idempotent No-op)
    res2 = client.post("/api/demo/seed")
    assert res2.status_code == 200
    data2 = res2.json()
    assert data2["project_id"] == proj_id

    # Crucial assertions: drawings count must remain EXACTLY 3 (no duplication)
    all_drawings_after = app_module.storage.list_drawings(proj_id)
    assert len(all_drawings_after) == 3, f"Expected 3 drawings, but found {len(all_drawings_after)}"

    # Crucial assertion: user reviews MUST NOT be overwritten or lost
    cands_after_res = client.get(f"/api/projects/{proj_id}/drawings/{img_dwg_id}/candidates")
    cands_after = cands_after_res.json()["candidates"]
    c_map = {c["candidate_id"]: c["status"] for c in cands_after}
    assert c_map[target_accept_id] == "ACCEPTED", "User ACCEPTED review decision was overwritten!"
    assert c_map[target_reject_id] == "REJECTED", "User REJECTED review decision was overwritten!"

    # 4. Explicit reset=True wipes and re-seeds clean
    res3 = client.post("/api/demo/seed?reset=true")
    assert res3.status_code == 200
    all_drawings_reset = app_module.storage.list_drawings(proj_id)
    assert len(all_drawings_reset) == 3
    cands_reset_res = client.get(f"/api/projects/{proj_id}/drawings/{img_dwg_id}/candidates")
    cands_reset = cands_reset_res.json()["candidates"]
    c_reset_map = {c["candidate_id"]: c["status"] for c in cands_reset}
    assert c_reset_map[target_accept_id] == "PROPOSED"


def test_visual_search_image(client: TestClient) -> None:
    """Verify visual template matching on raster image drawing finds all instances of the cropped symbol."""
    p_resp = client.post("/api/projects", json={"name": "Visual Search Image Test"})
    proj_id = p_resp.json()["id"]

    img_bytes = IMAGE_FIXTURE_PATH.read_bytes()
    up_resp = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("floorplan_sample.png", img_bytes, "image/png")},
    )
    assert up_resp.status_code == 200
    dwg_id = up_resp.json()["drawing_id"]

    # Crop one of the 4 lights at [227, 227, 273, 273]
    search_req = {
        "sample_bbox": [227.0, 227.0, 273.0, 273.0],
        "page_index": 0,
        "suggested_item_code": "LIGHT",
        "suggested_item_name": "다운라이트 조명",
        "threshold": 0.80,
    }
    s_resp = client.post(f"/api/projects/{proj_id}/drawings/{dwg_id}/search-symbol", json=search_req)
    assert s_resp.status_code == 200
    data = s_resp.json()

    assert data["drawing_id"] == dwg_id
    assert data["total_found"] == 4  # All 4 lights found
    assert len(data["candidates"]) == 4

    for c in data["candidates"]:
        assert c["suggested_item_code"] == "LIGHT"
        assert c["detection_source"] == "VISUAL_TEMPLATE"
        assert c["coordinate_space"] == "IMAGE_PIXEL"
        assert c["confidence"] >= 0.80
        assert "similarity_score" in c["raw_evidence"]
        assert c["status"] == "PROPOSED"

    # Verify candidates are persisted in storage
    all_cands = client.get(f"/api/projects/{proj_id}/drawings/{dwg_id}/candidates").json()["candidates"]
    assert any(c["detection_source"] == "VISUAL_TEMPLATE" for c in all_cands)


def test_visual_search_pdf(client: TestClient) -> None:
    """Verify visual template matching on vector PDF drawing finds rendered symbol in PDF points."""
    p_resp = client.post("/api/projects", json={"name": "Visual Search PDF Test"})
    proj_id = p_resp.json()["id"]

    pdf_bytes = PDF_FIXTURE_PATH.read_bytes()
    up_resp = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("pdf_c_vector_features.pdf", pdf_bytes, "application/pdf")},
    )
    assert up_resp.status_code == 200
    dwg_id = up_resp.json()["drawing_id"]

    # Search L-01 text/symbol patch on page 0
    search_req = {
        "sample_bbox": [95.0, 75.0, 130.0, 98.0],
        "page_index": 0,
        "suggested_item_code": "LIGHT",
        "suggested_item_name": "조명",
        "threshold": 0.75,
    }
    s_resp = client.post(f"/api/projects/{proj_id}/drawings/{dwg_id}/search-symbol", json=search_req)
    assert s_resp.status_code == 200
    data = s_resp.json()

    assert data["total_found"] >= 1
    first_cand = data["candidates"][0]
    assert first_cand["detection_source"] == "VISUAL_TEMPLATE"
    assert first_cand["coordinate_space"] == "PDF_PAGE_POINT"
    assert first_cand["page_index"] == 0
    assert first_cand["confidence"] >= 0.75


def test_visual_search_dxf(client: TestClient) -> None:
    """Verify visual template matching on DXF finds block occurrences overlapping selection."""
    p_resp = client.post("/api/projects", json={"name": "Visual Search DXF Test"})
    proj_id = p_resp.json()["id"]

    dxf_bytes = DXF_FIXTURE_PATH.read_bytes()
    up_resp = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("basic_devices.dxf", dxf_bytes, "application/dxf")},
    )
    assert up_resp.status_code == 200
    dwg_id = up_resp.json()["drawing_id"]

    # BBox overlapping LT-DL block at (1000, 2000)
    search_req = {
        "sample_bbox": [950.0, 1950.0, 1050.0, 2050.0],
        "suggested_item_code": "LIGHT",
        "suggested_item_name": "조명기구",
    }
    s_resp = client.post(f"/api/projects/{proj_id}/drawings/{dwg_id}/search-symbol", json=search_req)
    assert s_resp.status_code == 200
    data = s_resp.json()

    assert data["total_found"] == 6  # All 6 LT-DL blocks in drawing
    assert all(c["detection_source"] == "VISUAL_TEMPLATE" for c in data["candidates"])


def test_visual_search_validation_errors(client: TestClient) -> None:
    """Verify input validation: sample bbox too small, or non-existent drawing."""
    p_resp = client.post("/api/projects", json={"name": "Validation Test"})
    proj_id = p_resp.json()["id"]

    img_bytes = IMAGE_FIXTURE_PATH.read_bytes()
    up_resp = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("floorplan_sample.png", img_bytes, "image/png")},
    )
    dwg_id = up_resp.json()["drawing_id"]

    # 1. Bbox too small (width < 4px)
    too_small_req = {
        "sample_bbox": [10.0, 10.0, 12.0, 12.0],
        "suggested_item_code": "LIGHT",
    }
    bad_resp = client.post(f"/api/projects/{proj_id}/drawings/{dwg_id}/search-symbol", json=too_small_req)
    assert bad_resp.status_code == 400
    assert "너무 작습니다" in bad_resp.json()["detail"]

    # 2. Non-existent drawing
    not_found = client.post(
        f"/api/projects/{proj_id}/drawings/dwg_non_existent/search-symbol",
        json={"sample_bbox": [10.0, 10.0, 50.0, 50.0]},
    )
    assert not_found.status_code == 404


def test_p0_01_dxf_visual_search_no_double_counting_and_structural_blocking(client: TestClient) -> None:
    """Verify P0-01: DXF visual search candidates do NOT duplicate EA of mapped CAD blocks,

    and structurally blocked blocks fail-closed to 0 EA even if candidate is accepted.
    """
    from eai_api import app as app_module

    p_resp = client.post("/api/projects", json={"name": "P0-01 DXF Dedupe and Structural Block Test"})
    proj_id = p_resp.json()["id"]

    dxf_bytes = DXF_FIXTURE_PATH.read_bytes()
    up_resp = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("basic_devices.dxf", dxf_bytes, "application/dxf")},
    )
    dwg_id = up_resp.json()["drawing_id"]

    # 1. Map LT-DL to LIGHT (6 EA in basic_devices.dxf)
    rules = [
        ItemMappingRule(
            block_name="LT-DL",
            item_code="LIGHT",
            item_name="조명기구",
            unit="EA",
            specification="LED 50W",
        ).model_dump(mode="json")
    ]
    map_resp = client.put(f"/api/projects/{proj_id}/drawings/{dwg_id}/mappings", json=rules)
    assert map_resp.status_code == 200

    # Baseline takeoff before visual search: exactly 6 EA
    t0_resp = client.get(f"/api/projects/{proj_id}/drawings/{dwg_id}/takeoff")
    assert t0_resp.status_code == 200
    t0_data = t0_resp.json()
    light_row_0 = next(r for r in t0_data["summary"] if r["item_code"] == "LIGHT")
    assert light_row_0["quantity"] == 6

    # 2. Execute visual search on the LT-DL block
    search_req = {
        "sample_bbox": [950.0, 1950.0, 1050.0, 2050.0],
        "suggested_item_code": "LIGHT",
        "suggested_item_name": "조명기구",
    }
    s_resp = client.post(f"/api/projects/{proj_id}/drawings/{dwg_id}/search-symbol", json=search_req)
    assert s_resp.status_code == 200
    candidates = s_resp.json()["candidates"]
    assert len(candidates) == 6

    # 3. Bulk Accept all 6 visual candidates
    bulk_updates = [
        CandidateReviewUpdate(candidate_id=c["candidate_id"], status=CandidateStatus.ACCEPTED)
        for c in candidates
    ]
    rev_resp = client.put(
        f"/api/projects/{proj_id}/drawings/{dwg_id}/candidates/review",
        json=CandidateBulkReviewRequest(updates=bulk_updates).model_dump(mode="json"),
    )
    assert rev_resp.status_code == 200

    # 4. Critical P0-01 Assertion: Takeoff MUST STILL BE EXACTLY 6 EA, NOT 12 EA!
    t1_resp = client.get(f"/api/projects/{proj_id}/drawings/{dwg_id}/takeoff")
    assert t1_resp.status_code == 200
    t1_data = t1_resp.json()
    light_row_1 = next(r for r in t1_data["summary"] if r["item_code"] == "LIGHT")
    assert light_row_1["quantity"] == 6, f"Expected strictly 6 EA for LIGHT, got {light_row_1['quantity']} EA (double counting detected!)"

    # Evidence rows for LIGHT must be strictly 6 CAD insert rows
    light_evidence = [r for r in t1_data["evidence"] if r["mapped_item_code"] == "LIGHT" and r["inclusion_status"] in ("INCLUDED", "MANUAL_ADDED")]
    assert len(light_evidence) == 6, f"Expected strictly 6 included evidence rows, got {len(light_evidence)}"
    assert all(r["entity_type"] == "INSERT" for r in light_evidence), "Candidate rows must NOT duplicate mapped CAD INSERT rows"

    # 5. CSV Summary Assertion
    csv_sum_resp = client.get(f"/api/projects/{proj_id}/drawings/{dwg_id}/export/csv?type=summary")
    assert csv_sum_resp.status_code == 200
    csv_sum_reader = list(csv.reader(io.StringIO(csv_sum_resp.content.decode("utf-8-sig"))))
    sum_headers = csv_sum_reader[0]
    code_idx = sum_headers.index("item_code")
    qty_idx = sum_headers.index("quantity")
    light_csv_row = next(r for r in csv_sum_reader[1:] if r[code_idx] == "LIGHT")
    assert int(light_csv_row[qty_idx]) == 6, f"CSV Summary quantity must be strictly 6, got {light_csv_row[qty_idx]}"

    # 6. CSV Evidence Assertion
    csv_ev_resp = client.get(f"/api/projects/{proj_id}/drawings/{dwg_id}/export/csv?type=evidence")
    assert csv_ev_resp.status_code == 200
    csv_ev_reader = list(csv.reader(io.StringIO(csv_ev_resp.content.decode("utf-8-sig"))))
    ev_headers = csv_ev_reader[0]
    ev_code_idx = ev_headers.index("mapped_item_code")
    ev_inc_idx = ev_headers.index("inclusion_status")
    light_ev_csv_rows = [r for r in csv_ev_reader[1:] if r[ev_code_idx] == "LIGHT" and r[ev_inc_idx] == "INCLUDED"]
    assert len(light_ev_csv_rows) == 6, f"CSV Evidence count must be strictly 6, got {len(light_ev_csv_rows)}"

    # 7. Excel Openpyxl Actual Cell Value Assertion
    xlsx_resp = client.get(f"/api/projects/{proj_id}/drawings/{dwg_id}/export/xlsx")
    assert xlsx_resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_resp.content))
    assert "Takeoff Summary" in wb.sheetnames
    ws_sum = wb["Takeoff Summary"]
    headers = [str(cell.value) for cell in ws_sum[1]]
    code_col = next(i + 1 for i, h in enumerate(headers) if "item_code" in h)
    qty_col = next(i + 1 for i, h in enumerate(headers) if "quantity" in h)

    found_xlsx_light = False
    for row in range(2, ws_sum.max_row + 1):
        code_val = ws_sum.cell(row=row, column=code_col).value
        qty_val = ws_sum.cell(row=row, column=qty_col).value
        if code_val == "LIGHT":
            assert int(qty_val) == 6, f"Excel Summary cell value must be strictly 6, got {qty_val}"
            found_xlsx_light = True
            break
    assert found_xlsx_light, "LIGHT row not found in XLSX Takeoff Summary"

    # 8. Structural Block Fail-Closed Verification (P0-01)
    # Inject a structurally blocked candidate on this drawing and accept it
    blocked_cand = SymbolCandidate(
        candidate_id=f"cand_blocked_dxf_{dwg_id}_1",
        drawing_id=dwg_id,
        page_index=0,
        coordinate_space="WORLD",
        x=500.0,
        y=500.0,
        bbox=[485.0, 485.0, 515.0, 515.0],
        suggested_item_code="PANEL",
        suggested_item_name="분전반 (검수불가: 동적/익명 블록)",
        confidence=0.0,
        detection_source="DXF_BLOCK",
        status=CandidateStatus.ACCEPTED,
        raw_evidence={
            "is_structurally_blocked": True,
            "structural_block_reason": "동적/익명 블록: 가시성·구조 검수 필요",
        },
    )
    existing_cands = app_module.storage.get_candidates(proj_id, dwg_id)
    app_module.storage.save_candidates(proj_id, dwg_id, [*existing_cands, blocked_cand])

    # Recompute takeoff: structurally blocked candidate MUST NEVER produce quantity (must be 0 EA)
    t2_resp = client.get(f"/api/projects/{proj_id}/drawings/{dwg_id}/takeoff")
    t2_data = t2_resp.json()
    panel_rows = [r for r in t2_data["summary"] if r["item_code"] == "PANEL"]
    assert len(panel_rows) == 0 or panel_rows[0]["quantity"] == 0, "Structurally blocked candidate was illegally revived into positive EA!"


def test_p0_02_spatial_deduplication_and_review_preservation(client: TestClient) -> None:
    """Verify P0-02: Repeated or cross searches on the same physical symbol location

    do NOT produce duplicate candidates or multiply EA, and strictly preserve user review decisions.
    Different PDF pages remain separate.
    """
    from eai_api import app as app_module

    p_resp = client.post("/api/projects", json={"name": "P0-02 Spatial Dedupe Test"})
    proj_id = p_resp.json()["id"]

    img_bytes = IMAGE_FIXTURE_PATH.read_bytes()
    up_resp = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("floorplan_sample.png", img_bytes, "image/png")},
    )
    dwg_id = up_resp.json()["drawing_id"]

    # 1. Search #1 targeting a light symbol around (250, 250)
    search_req_1 = {
        "sample_bbox": [227.0, 227.0, 273.0, 273.0],
        "suggested_item_code": "LIGHT",
        "suggested_item_name": "조명기구",
        "threshold": 0.75,
    }
    s1_resp = client.post(f"/api/projects/{proj_id}/drawings/{dwg_id}/search-symbol", json=search_req_1)
    assert s1_resp.status_code == 200
    cands_1 = s1_resp.json()["candidates"]
    assert len(cands_1) >= 1
    total_found_1 = len(cands_1)

    # 2. User explicitly reviews candidate #0 as ACCEPTED
    target_cand_id = cands_1[0]["candidate_id"]
    rev_resp = client.put(
        f"/api/projects/{proj_id}/drawings/{dwg_id}/candidates/review",
        json=CandidateBulkReviewRequest(
            updates=[
                CandidateReviewUpdate(candidate_id=target_cand_id, status=CandidateStatus.ACCEPTED, review_note="사용자 수동 확정")
            ]
        ).model_dump(mode="json"),
    )
    assert rev_resp.status_code == 200

    # 3. Search #2 on the same image with slightly shifted bbox and different threshold
    search_req_2 = {
        "sample_bbox": [225.0, 225.0, 275.0, 275.0],
        "suggested_item_code": "LIGHT",
        "suggested_item_name": "조명기구",
        "threshold": 0.70,
    }
    s2_resp = client.post(f"/api/projects/{proj_id}/drawings/{dwg_id}/search-symbol", json=search_req_2)
    assert s2_resp.status_code == 200

    # 4. Check stored candidates in storage: MUST NOT duplicate the candidate count at same locations
    stored_cands = app_module.storage.get_candidates(proj_id, dwg_id)
    # The count should match total physical locations, NOT double!
    assert len(stored_cands) == total_found_1, f"Expected {total_found_1} candidates, got {len(stored_cands)} (duplication occurred!)"

    # User review decision on candidate #0 MUST BE STRICTLY PRESERVED!
    target_stored = next(c for c in stored_cands if c.candidate_id == target_cand_id)
    assert target_stored.status == CandidateStatus.ACCEPTED, "User review decision was reset to PROPOSED on re-search!"
    assert target_stored.review_note == "사용자 수동 확정"

    # 5. Takeoff assertion: one physical location produces AT MOST 1 EA
    # Accept all stored candidates
    bulk_all = [CandidateReviewUpdate(candidate_id=c.candidate_id, status=CandidateStatus.ACCEPTED) for c in stored_cands]
    client.put(
        f"/api/projects/{proj_id}/drawings/{dwg_id}/candidates/review",
        json=CandidateBulkReviewRequest(updates=bulk_all).model_dump(mode="json"),
    )
    t_resp = client.get(f"/api/projects/{proj_id}/drawings/{dwg_id}/takeoff")
    t_data = t_resp.json()
    light_row = next(r for r in t_data["summary"] if r["item_code"] == "LIGHT")
    # Must equal total physical lights, NOT duplicated
    assert light_row["quantity"] == total_found_1

    # 6. Multi-page PDF independence: same coordinates on Page 0 and Page 1 are separate physical locations
    cand_p0 = SymbolCandidate(
        candidate_id="pdf_cand_p0",
        drawing_id="dwg_pdf_test",
        page_index=0,
        coordinate_space="PDF_PAGE_POINT",
        x=150.0,
        y=150.0,
        bbox=[140.0, 140.0, 160.0, 160.0],
        suggested_item_code="LIGHT",
        suggested_item_name="조명기구",
        confidence=0.9,
        detection_source="VISUAL_TEMPLATE",
    )
    cand_p1 = SymbolCandidate(
        candidate_id="pdf_cand_p1",
        drawing_id="dwg_pdf_test",
        page_index=1,
        coordinate_space="PDF_PAGE_POINT",
        x=150.0,
        y=150.0,
        bbox=[140.0, 140.0, 160.0, 160.0],
        suggested_item_code="LIGHT",
        suggested_item_name="조명기구",
        confidence=0.9,
        detection_source="VISUAL_TEMPLATE",
    )
    from eai_api.visual_search import are_candidates_spatially_coincident
    assert not are_candidates_spatially_coincident(cand_p0, cand_p1), "Different PDF pages must NOT be considered coincident!"


def test_p0_review_01_reviewed_candidates_immutable_on_research(client: TestClient) -> None:
    """[P0-REVIEW-01] User-reviewed candidates (ACCEPTED/REJECTED) must never be silently reclassified

    or have their effective item codes/review notes/timestamps mutated by a subsequent search
    targeting the same location with a different item_code.
    """
    from eai_api import app as app_module
    from eai_api.visual_search import merge_candidates_spatial

    p_resp = client.post("/api/projects", json={"name": "P0-REVIEW-01 Reclassification Defense Test"})
    proj_id = p_resp.json()["id"]

    img_bytes = IMAGE_FIXTURE_PATH.read_bytes()
    up_resp = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("floorplan_sample.png", img_bytes, "image/png")},
    )
    dwg_id = up_resp.json()["drawing_id"]

    # 1. Search #1 targeting LIGHT around (250, 250)
    search_req_1 = {
        "sample_bbox": [227.0, 227.0, 273.0, 273.0],
        "suggested_item_code": "LIGHT",
        "suggested_item_name": "조명기구",
        "threshold": 0.75,
    }
    s1_resp = client.post(f"/api/projects/{proj_id}/drawings/{dwg_id}/search-symbol", json=search_req_1)
    assert s1_resp.status_code == 200
    cands_1 = s1_resp.json()["candidates"]
    assert len(cands_1) >= 1
    target_cand_id = cands_1[0]["candidate_id"]

    # 2. User explicitly ACCEPTS candidate (reviewed_item_code=None, accepting suggested LIGHT)
    review_note = "조명기구 정규 승인 완료"
    rev_resp = client.put(
        f"/api/projects/{proj_id}/drawings/{dwg_id}/candidates/review",
        json=CandidateBulkReviewRequest(
            updates=[
                CandidateReviewUpdate(
                    candidate_id=target_cand_id,
                    status=CandidateStatus.ACCEPTED,
                    reviewed_item_code=None,  # Explicitly None to test the default acceptance path
                    review_note=review_note,
                )
            ]
        ).model_dump(mode="json"),
    )
    assert rev_resp.status_code == 200

    # Verify reviewed candidate state in storage before re-search
    stored_before = app_module.storage.get_candidates(proj_id, dwg_id)
    target_before = next(c for c in stored_before if c.candidate_id == target_cand_id)
    assert target_before.status == CandidateStatus.ACCEPTED
    assert target_before.suggested_item_code == "LIGHT"
    assert target_before.reviewed_item_code is None
    assert target_before.reviewed_at is not None
    assert target_before.review_note == review_note
    original_reviewed_at = target_before.reviewed_at

    # 3. Search #2 on the EXACT same location with suggested_item_code="OUTLET"
    search_req_2 = {
        "sample_bbox": [227.0, 227.0, 273.0, 273.0],
        "suggested_item_code": "OUTLET",
        "suggested_item_name": "전열콘센트",
        "threshold": 0.75,
    }
    s2_resp = client.post(f"/api/projects/{proj_id}/drawings/{dwg_id}/search-symbol", json=search_req_2)
    assert s2_resp.status_code == 200

    # 4. Check candidate in storage: MUST NOT be reclassified to OUTLET!
    stored_after = app_module.storage.get_candidates(proj_id, dwg_id)
    target_after = next(c for c in stored_after if c.candidate_id == target_cand_id)
    assert target_after.status == CandidateStatus.ACCEPTED, "Review status was lost!"
    assert target_after.suggested_item_code == "LIGHT", f"Candidate suggested_item_code was illegally overridden to {target_after.suggested_item_code}!"
    assert target_after.suggested_item_name == "조명기구", "Candidate suggested_item_name was mutated!"
    assert target_after.reviewed_item_code is None, "reviewed_item_code was unexpectedly mutated!"
    assert target_after.reviewed_at == original_reviewed_at, "reviewed_at was reset or cleared!"
    assert target_after.review_note == review_note, "review_note was lost!"

    # 5. Takeoff assertion: strictly LIGHT 1 EA, OUTLET 0 EA
    t_resp = client.get(f"/api/projects/{proj_id}/drawings/{dwg_id}/takeoff")
    assert t_resp.status_code == 200
    t_data = t_resp.json()
    summary_map = {r["item_code"]: r["quantity"] for r in t_data["summary"]}
    assert summary_map.get("LIGHT") == 1, f"Expected LIGHT 1 EA, got {summary_map.get('LIGHT')}"
    assert summary_map.get("OUTLET", 0) == 0, f"OUTLET must be 0 EA, got {summary_map.get('OUTLET')}"

    # 6. CSV Export assertions
    csv_sum_resp = client.get(f"/api/projects/{proj_id}/drawings/{dwg_id}/export/csv?type=summary")
    assert csv_sum_resp.status_code == 200
    csv_rows = list(csv.reader(io.StringIO(csv_sum_resp.content.decode("utf-8-sig"))))
    header = csv_rows[0]
    item_code_idx = header.index("item_code")
    qty_idx = header.index("quantity")
    light_csv = next((r for r in csv_rows[1:] if r[item_code_idx] == "LIGHT"), None)
    outlet_csv = next((r for r in csv_rows[1:] if r[item_code_idx] == "OUTLET"), None)
    assert light_csv is not None and int(light_csv[qty_idx]) == 1
    assert outlet_csv is None

    # 7. Excel (.xlsx) openpyxl exact cell value verification
    xlsx_resp = client.get(f"/api/projects/{proj_id}/drawings/{dwg_id}/export/xlsx")
    assert xlsx_resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_resp.content))
    assert "Takeoff Summary" in wb.sheetnames
    ws = wb["Takeoff Summary"]
    headers = [str(cell.value) for cell in ws[1]]
    code_col = next(i + 1 for i, h in enumerate(headers) if "item_code" in h)
    qty_col = next(i + 1 for i, h in enumerate(headers) if "quantity" in h)
    found_light_qty: int | None = None
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=code_col).value == "LIGHT":
            found_light_qty = int(ws.cell(row=row, column=qty_col).value or 0)
    assert found_light_qty == 1, f"Excel Takeoff Summary LIGHT cell must be 1, got {found_light_qty}"

    # 8. REJECTED candidate preservation
    reject_cand = SymbolCandidate(
        candidate_id="cand_rej_test",
        drawing_id=dwg_id,
        page_index=0,
        coordinate_space="IMAGE_PIXEL",
        x=500.0,
        y=500.0,
        bbox=[490.0, 490.0, 510.0, 510.0],
        suggested_item_code="SWITCH",
        suggested_item_name="스위치",
        confidence=0.88,
        detection_source="IMAGE_TEMPLATE",
        status=CandidateStatus.REJECTED,
        review_note="노이즈 오탐 반려",
    )
    incoming_switch = SymbolCandidate(
        candidate_id="cand_new_switch",
        drawing_id=dwg_id,
        page_index=0,
        coordinate_space="IMAGE_PIXEL",
        x=500.0,
        y=500.0,
        bbox=[490.0, 490.0, 510.0, 510.0],
        suggested_item_code="PANEL",
        suggested_item_name="분전반",
        confidence=0.92,
        detection_source="IMAGE_TEMPLATE",
        status=CandidateStatus.PROPOSED,
    )
    merged_rej = merge_candidates_spatial([reject_cand], [incoming_switch])
    assert len(merged_rej) == 1
    assert merged_rej[0].status == CandidateStatus.REJECTED
    assert merged_rej[0].suggested_item_code == "SWITCH"
    assert merged_rej[0].review_note == "노이즈 오탐 반려"


def test_p0_review_02_adjacent_close_symbols_not_eliminated(client: TestClient) -> None:
    """[P0-REVIEW-02] Closely adjacent, non-overlapping physical equipment must NOT be merged

    by an arbitrary distance radius (e.g. 15px). Disjoint bounding boxes must remain separate
    candidates and produce exact multi-EA counts upon acceptance.
    """
    from eai_api import app as app_module
    from eai_api.visual_search import are_candidates_spatially_coincident, merge_candidates_spatial

    # 1. Two small non-overlapping close equipment on the same IMAGE drawing
    # c1: bbox [0, 0, 4, 4], center (2, 2)
    # c2: bbox [10, 0, 14, 4], center (12, 2)
    # Center distance is 10.0 <= 15.0, but bounding boxes are strictly disjoint (inter == 0.0)
    c1 = SymbolCandidate(
        candidate_id="cand_adj_1",
        drawing_id="dwg_adj_test",
        page_index=0,
        coordinate_space="IMAGE_PIXEL",
        x=2.0,
        y=2.0,
        bbox=[0.0, 0.0, 4.0, 4.0],
        suggested_item_code="LIGHT",
        suggested_item_name="소형 조명 1",
        confidence=0.9,
        detection_source="IMAGE_TEMPLATE",
        status=CandidateStatus.PROPOSED,
    )
    c2 = SymbolCandidate(
        candidate_id="cand_adj_2",
        drawing_id="dwg_adj_test",
        page_index=0,
        coordinate_space="IMAGE_PIXEL",
        x=12.0,
        y=2.0,
        bbox=[10.0, 0.0, 14.0, 4.0],
        suggested_item_code="LIGHT",
        suggested_item_name="소형 조명 2",
        confidence=0.9,
        detection_source="IMAGE_TEMPLATE",
        status=CandidateStatus.PROPOSED,
    )

    # Invariant: Disjoint bounding boxes are NEVER spatially coincident
    assert not are_candidates_spatially_coincident(c1, c2), "Non-overlapping close symbols must NOT be considered coincident!"

    # Invariant: Merging them keeps BOTH candidates
    merged = merge_candidates_spatial([c1], [c2])
    assert len(merged) == 2, f"Expected 2 candidates, got {len(merged)} (under-counting occurred!)"

    # 2. Full Takeoff and CSV validation on project with adjacent symbols
    p_resp = client.post("/api/projects", json={"name": "Adjacent Symbols EA Test"})
    proj_id = p_resp.json()["id"]

    img_bytes = IMAGE_FIXTURE_PATH.read_bytes()
    up_resp = client.post(
        f"/api/projects/{proj_id}/drawings",
        files={"file": ("floorplan_sample.png", img_bytes, "image/png")},
    )
    dwg_id = up_resp.json()["drawing_id"]

    # Save both candidates as ACCEPTED
    c1_dwg = c1.model_copy(update={"drawing_id": dwg_id, "status": CandidateStatus.ACCEPTED})
    c2_dwg = c2.model_copy(update={"drawing_id": dwg_id, "status": CandidateStatus.ACCEPTED})
    app_module.storage.save_candidates(proj_id, dwg_id, [c1_dwg, c2_dwg])

    # Calculate Takeoff: MUST strictly be 2 EA (not reduced to 1 EA!)
    t_resp = client.get(f"/api/projects/{proj_id}/drawings/{dwg_id}/takeoff")
    assert t_resp.status_code == 200
    t_data = t_resp.json()
    light_row = next((r for r in t_data["summary"] if r["item_code"] == "LIGHT"), None)
    assert light_row is not None
    assert light_row["quantity"] == 2, f"Adjacent symbols must produce 2 EA, got {light_row['quantity']}"

    # Verify Evidence rows count is strictly 2
    ev_cands = [e for e in t_data["evidence"] if e["entity_type"] == "SYMBOL_CANDIDATE"]
    assert len(ev_cands) == 2, f"Expected 2 evidence rows, got {len(ev_cands)}"

    # 3. DXF with different source_object_id: NEVER coincident even if bounding boxes overlap
    dxf_c1 = SymbolCandidate(
        candidate_id="dxf_cand_1",
        drawing_id="dwg_dxf_test",
        page_index=0,
        coordinate_space="WORLD",
        x=100.0,
        y=100.0,
        bbox=[85.0, 85.0, 115.0, 115.0],
        suggested_item_code="LIGHT",
        suggested_item_name="조명 1",
        confidence=0.95,
        detection_source="DXF_BLOCK",
        raw_evidence={"source_object_id": "cad_obj_001"},
    )
    dxf_c2 = SymbolCandidate(
        candidate_id="dxf_cand_2",
        drawing_id="dwg_dxf_test",
        page_index=0,
        coordinate_space="WORLD",
        x=105.0,
        y=105.0,
        bbox=[90.0, 90.0, 120.0, 120.0],
        suggested_item_code="LIGHT",
        suggested_item_name="조명 2",
        confidence=0.95,
        detection_source="DXF_BLOCK",
        raw_evidence={"source_object_id": "cad_obj_002"},
    )
    assert not are_candidates_spatially_coincident(dxf_c1, dxf_c2), "Different CAD source_object_ids must NEVER be coincident!"
    merged_dxf = merge_candidates_spatial([dxf_c1], [dxf_c2])
    assert len(merged_dxf) == 2

